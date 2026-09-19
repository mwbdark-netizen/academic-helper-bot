# -*- coding: utf-8 -*-
"""
المنصة الأكاديمية الذكية - بوت تلغرام
المرحلة الثانية - علوم الحاسوب

الخدمات: تقرير (Word / PDF) - شرح موضوع أو صفحة ملزمة - حل مسائل
         تدقيق أكواد - بنك أسئلة
لغة الناتج تتبع لغة دراسة المادة، ما عدا الشرح فيكون بالعربية افتراضياً
مع الحفاظ على المصطلحات الإنكليزية. الصور مدعومة في المسائل والأكواد والشرح.

إحصائيات دائمة: عدد المستخدمين ومرات استخدام كل خدمة، مع نسخ احتياطي
تلقائي إلى قناة خاصة واستعادة تلقائية بعد كل إعادة نشر.

توليد PDF: خط يونيكودي واحد لجميع الأنماط، كشف اتجاه كل فقرة تلقائياً
(عربية أو لاتينية) داخل الملف نفسه، وتنظيف الرموز غير المدعومة، وكتل
الكود تُخرج بخط ثابت العرض مع الحفاظ على الإزاحة.

اقتراح يوتيوب: عبارات بحث دقيقة مبنية على موضوع الصفحة واسم المادة،
استبعاد المحتوى المدرسي، ترتيب النتائج بثلاثة عوامل (تطابق الموضوع،
تطابق المادة، ترتيب العبارة)، فكّ ترميز العناوين، وذاكرة مؤقتة
تقلّل استهلاك الحصة اليومية.
"""

import os
import io
import re
import json
import base64
import logging
import asyncio
import threading
from html import escape
import html as _h          # لفكّ ترميز عناوين يوتيوب؛ لا يتعارض مع escape أعلاه
from urllib.parse import quote_plus
from datetime import datetime, timezone, timedelta
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler

import httpx
from groq import AsyncGroq

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:
    from PIL import Image
    PIL_OK = True
except Exception:
    PIL_OK = False

try:
    from fpdf import FPDF
    FPDF_OK = True
except Exception:
    FPDF_OK = False

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    BIDI_OK = True
except Exception:
    BIDI_OK = False

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.constants import ChatAction
from telegram.error import Forbidden, RetryAfter
from telegram.ext import (
    Application,
    CommandHandler,
    CallbackQueryHandler,
    MessageHandler,
    ContextTypes,
    PicklePersistence,
    filters,
)

try:
    from telegram import LinkPreviewOptions
    NO_PREVIEW = LinkPreviewOptions(is_disabled=True)
except Exception:
    NO_PREVIEW = None

# ══════════════════════════════ الإعدادات ══════════════════════════════

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("telegram.ext.Application").setLevel(logging.WARNING)
logging.getLogger("fontTools").setLevel(logging.ERROR)
logger = logging.getLogger("academic_bot")

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "").strip()

_admin_raw = os.getenv("ADMIN_ID", "0").strip()
ADMIN_ID = int(_admin_raw) if _admin_raw.lstrip("-").isdigit() else 0

PORT = int(os.getenv("PORT", "10000"))

DATA_DIR = (os.getenv("DATA_DIR", ".").rstrip("/") or ".")
STATS_FILE = os.path.join(DATA_DIR, "bot_stats.json")
PERSIST_FILE = os.path.join(DATA_DIR, "bot_persistence.pkl")

# قناة خاصة لحفظ نسخة الإحصائيات (تبدأ بـ -100)
BACKUP_CHAT_ID = os.getenv("BACKUP_CHAT_ID", "").strip()
BACKUP_HOURS = float(os.getenv("BACKUP_HOURS", "6") or 6)
BAGHDAD_TZ = timezone(timedelta(hours=3))
KEEP_DAYS = 60

# بيانات الغلاف الاختياري
UNIVERSITY = os.getenv("UNIVERSITY", "جامعة القادسية")
COLLEGE = os.getenv("COLLEGE", "كلية علوم الحاسوب وتكنولوجيا المعلومات")
DEPARTMENT = os.getenv("DEPARTMENT", "قسم علوم الحاسوب")
STAGE = os.getenv("STAGE", "المرحلة الثانية")

# السنة الدراسية: تُحسب تلقائياً حسب الشهر، أو تُثبَّت عبر المتغيّر ACADEMIC_YEAR
ACADEMIC_YEAR_ENV = os.getenv("ACADEMIC_YEAR", "").strip()
ACADEMIC_START_MONTH = int(os.getenv("ACADEMIC_START_MONTH", "8") or 8)


def academic_year() -> str:
    """بعد شهر آب تبدأ سنة دراسية جديدة: 2026 - 2027 مثلاً."""
    if ACADEMIC_YEAR_ENV:
        return ACADEMIC_YEAR_ENV
    now = datetime.now(BAGHDAD_TZ)
    y = now.year if now.month >= ACADEMIC_START_MONTH else now.year - 1
    return f"{y} - {y + 1}"


# الخطوط
ARABIC_FONT = os.getenv("ARABIC_FONT", "Arial")
LATIN_FONT = os.getenv("LATIN_FONT", "Times New Roman")
MONO_FONT = os.getenv("MONO_FONT", "Courier New")   # خط كتل الكود في Word
FONT_DIR = os.getenv("FONT_DIR", "fonts")
AR_TTF = os.path.join(FONT_DIR, os.getenv("AR_TTF", "Amiri-Regular.ttf"))
AR_TTF_BOLD = os.path.join(FONT_DIR, os.getenv("AR_TTF_BOLD", "Amiri-Bold.ttf"))
# خط ثابت العرض للكود في PDF؛ اختياري، وإن غاب يُستخدم الخط العادي
MONO_TTF = os.path.join(FONT_DIR, os.getenv("MONO_TTF", "DejaVuSansMono.ttf"))

MAX_INPUT_CHARS = 8000
TELEGRAM_CHUNK = 3800

# حدود الصور
MAX_IMAGES = 5
IMG_MAX_SIDE = 1600
IMG_QUALITY = 85
IMG_MAX_BYTES = 4 * 1024 * 1024
ALBUM_WAIT = 2.5

# يوتيوب
YOUTUBE_API_KEY = os.getenv("YOUTUBE_API_KEY", "").strip()
YT_RESULTS = 3
YT_FETCH = 8          # نطلب نتائج أكثر ثم نفلترها؛ الكلفة 100 وحدة سواء 3 أو 8
YT_MAX_QUERIES = int(os.getenv("YT_MAX_QUERIES", "2") or 2)   # كل عبارة = 100 وحدة
YT_TIMEOUT = 12.0
YT_API = "https://www.googleapis.com/youtube/v3/search"
YT_SEARCH_URL = "https://www.youtube.com/results?search_query="

# لغة الشرح: عربية افتراضياً مع الاحتفاظ بالمصطلحات الإنكليزية
EXPLAIN_IN_ARABIC = os.getenv("EXPLAIN_IN_ARABIC", "1").strip() != "0"

BUSY_USERS: set[int] = set()

# ══════════════════════════════ المواد ══════════════════════════════
# lang: لغة المادة -> لغة الناتج      is_code: تظهر في خدمة تدقيق الأكواد

SUBJECTS = {
    # ── الكورس الأول ──
    "t1_oop":  {"name": "البرمجة الكيانية (OOP)",
                "term": 1, "lang": "en", "is_code": True,
                "en": "Object Oriented Programming"},
    "t1_mp":   {"name": "المعالجات الدقيقة (Microprocessors)",
                "term": 1, "lang": "en", "is_code": True,
                "en": "Microprocessors"},
    "t1_num":  {"name": "الطرق العددية (Numerical Methods)",
                "term": 1, "lang": "en", "is_code": True,
                "en": "Numerical Methods"},
    "t1_algo": {"name": "تحليل وتصميم الخوارزميات (Algorithms)",
                "term": 1, "lang": "en", "is_code": True,
                "en": "Analysis and Design of Algorithms"},
    "t1_emg":  {"name": "الإدارة الإلكترونية",
                "term": 1, "lang": "ar", "is_code": False,
                "en": "Electronic Management"},
    "t1_ar":   {"name": "اللغة العربية",
                "term": 1, "lang": "ar", "is_code": False,
                "en": "Arabic Language"},
    # ── الكورس الثاني ──
    "t2_ds":   {"name": "هياكل البيانات (Data Structures)",
                "term": 2, "lang": "en", "is_code": True,
                "en": "Data Structures"},
    "t2_java": {"name": "جافا (Java)",
                "term": 2, "lang": "en", "is_code": True,
                "en": "Java Programming"},
    "t2_toc":  {"name": "النظرية الاحتسابية (Theory of Computation)",
                "term": 2, "lang": "en", "is_code": False,
                "en": "Theory of Computation"},
    "t2_stat": {"name": "الإحصاء والاحتمالية (Probability & Statistics)",
                "term": 2, "lang": "en", "is_code": False,
                "en": "Probability and Statistics"},
    "t2_arch": {"name": "معمارية الحاسوب (Computer Architecture)",
                "term": 2, "lang": "en", "is_code": True,
                "en": "Computer Architecture"},
    "t2_en":   {"name": "اللغة الإنكليزية",
                "term": 2, "lang": "en", "is_code": False,
                "en": "English Language"},
}

SERVICE_LABELS = {
    "rep": "تقرير أكاديمي",
    "exp": "شرح تعليمي",
    "math": "حل مسألة",
    "code": "تدقيق كود",
    "quiz": "بنك أسئلة",
    "exp_img": "شرح من صورة",
    "math_img": "مسألة من صورة",
    "code_img": "كود من صورة",
}

# ══════════════════════════════ نماذج Groq ══════════════════════════════

TEXT_CHAIN = [
    {"id": "openai/gpt-oss-120b", "max_tokens": 8192, "extra": {"include_reasoning": False}},
    {"id": "openai/gpt-oss-20b",  "max_tokens": 8192, "extra": {"include_reasoning": False}},
    {"id": "qwen/qwen3.6-27b",    "max_tokens": 8192, "extra": {"reasoning_format": "hidden"}},
]

VISION_CHAIN = [
    {"id": "qwen/qwen3.6-27b", "max_images": 5, "max_tokens": 8192,
     "extra": {"reasoning_format": "hidden"}},
    {"id": "qwen/qwen3.8-27b", "max_images": 3, "max_tokens": 8192,
     "extra": {"reasoning_format": "hidden"}},
]

_groq_client: AsyncGroq | None = None


def get_groq_client() -> AsyncGroq:
    global _groq_client
    if _groq_client is None:
        if not GROQ_API_KEY:
            raise RuntimeError("GROQ_API_KEY غير موجود في متغيرات البيئة.")
        _groq_client = AsyncGroq(api_key=GROQ_API_KEY, timeout=90.0, max_retries=1)
    return _groq_client


def _build_messages(system: str, user: str, merge_system: bool):
    if merge_system:
        return [{"role": "user", "content": f"{system}\n\n---\n\n{user}"}]
    return [
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ]


async def query_ai_text(system: str, user: str, merge_system: bool = False,
                        max_tokens: int | None = None, temperature: float = 0.3) -> str:
    client = get_groq_client()
    last_err: Exception | None = None
    for m in TEXT_CHAIN:
        base = dict(
            model=m["id"],
            messages=_build_messages(system, user, merge_system),
            temperature=temperature,
            max_completion_tokens=max_tokens or m["max_tokens"],
            top_p=0.95,
        )
        for attempt in (1, 2):
            kwargs = dict(base)
            if attempt == 1:
                kwargs.update(m.get("extra", {}))
            try:
                resp = await client.chat.completions.create(**kwargs)
                txt = (resp.choices[0].message.content or "").strip()
                if txt:
                    return txt
                last_err = RuntimeError("رد فارغ")
            except Exception as e:
                last_err = e
                logger.warning("text model %s (attempt %s) failed: %s", m["id"], attempt, e)
    raise RuntimeError(f"تعذّر الحصول على رد من جميع النماذج: {last_err}")


class VisionUnavailable(RuntimeError):
    """سقطت سلسلة نماذج الرؤية كلها؛ نماذج الرؤية على Groq من فئة المعاينة."""


async def query_ai_vision(prompt: str, images: list[str]) -> str:
    client = get_groq_client()
    last_err: Exception | None = None
    for m in VISION_CHAIN:
        content = [{"type": "text", "text": prompt}]
        for b64 in images[: m["max_images"]]:
            content.append({
                "type": "image_url",
                "image_url": {"url": f"data:image/jpeg;base64,{b64}"},
            })
        base = dict(
            model=m["id"],
            messages=[{"role": "user", "content": content}],
            temperature=0.2,
            max_completion_tokens=m["max_tokens"],
        )
        for attempt in (1, 2):
            kwargs = dict(base)
            if attempt == 1:
                kwargs.update(m.get("extra", {}))
            try:
                resp = await client.chat.completions.create(**kwargs)
                txt = (resp.choices[0].message.content or "").strip()
                if txt:
                    return txt
                last_err = RuntimeError("رد فارغ")
            except Exception as e:
                last_err = e
                logger.warning("vision model %s (attempt %s) failed: %s", m["id"], attempt, e)
    raise VisionUnavailable(f"تعذّر تحليل الصورة: {last_err}")


# ══════════════════════════════ يوتيوب ══════════════════════════════

_YT_CACHE: dict[str, list[dict]] = {}
_YT_CACHE_MAX = 300

_TOPIC_NOISE = re.compile(
    r"^\s*(?:الموضوع|العنوان|شرح|مقدمة\s+(?:في|عن)|مراجعة|ملخص|درس|محاضرة|"
    r"topic|title|introduction\s+to|overview\s+of|lecture)\s*[:：\-–]?\s+",
    re.IGNORECASE)

# مؤشرات المحتوى المدرسي: تُستبعد لأن الطالب جامعي
# أُضيفت صيغ البكالوريا المغاربية والسادس العلمي بعد ظهور «تانية بكالوريا» في النتائج
_SCHOOL_NOISE = re.compile(
    r"(الصف\s*(?:الأول|الثاني|الثالث|الرابع|الخامس|السادس|\d)|"
    r"الثانوي|ثانوي|المتوسط|الابتدائي|ابتدائي|الإعدادي|اعدادي|"
    r"بكالوريا|باكالوريا|\bباك\b|تانية\s*باك|أولى\s*باك|"
    r"السادس\s*العلمي|السادس\s*الأدبي|"
    r"المنهج\s*السعودي|المنهج\s*المصري|مهارات\s*رقمية|تقنية\s*رقمية|"
    r"حاسب\s*[1-6]\b|للأطفال|"
    r"grade\s*\d|\b\d{1,2}(?:th|st|nd|rd)\s+grade\b|"
    r"high\s*school|middle\s*school|for\s+kids)",
    re.IGNORECASE)

# محتوى تعريفي للمبتدئين: صحيح لكنه أدنى من المستوى الجامعي، فيُخصم لا يُستبعد
_BEGINNER = re.compile(
    r"(for\s+beginners|beginners?\s+guide|basics\s+for|simple\s+explanation|"
    r"animated|teardown|crash\s*course|explained\s+in\s+\d+|"
    r"للمبتدئين|للمبتدئ|بأبسط\s*الطرق)",
    re.IGNORECASE)

# مواضيع مجاورة: صحيحة أكاديمياً لكنها ليست موضوع الصفحة
_NEAR_TOPIC = re.compile(
    r"(operating\s*system|أنظمة\s*التشغيل|نظم\s*التشغيل|compiler|مترجم|"
    r"شبكات|network|database|قواعد\s*البيانات|assembly)", re.IGNORECASE)

# مصطلح إنكليزي بين قوسين داخل عنوان الموضوع
_PAREN_EN = re.compile(r"[\(\[]\s*([A-Za-z][A-Za-z0-9 &/\.\-]{2,40})\s*[\)\]]")


def _norm_w(w: str) -> str:
    """تطبيع الكلمة: حذف التشكيل وتوحيد الهمزات وإسقاط أداة التعريف."""
    w = re.sub(r"[\u064B-\u0652\u0640]", "", (w or "").lower())
    w = (w.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
          .replace("ى", "ي").replace("ة", "ه"))
    return w[2:] if w.startswith("ال") and len(w) > 4 else w


# تُطبَّع عند التعريف، وإلا لن تطابق ناتج _norm_w إطلاقاً (محاضرة ← محاضره)
_STOP_W = {_norm_w(w) for w in (
    "شرح", "محاضرة", "محاضرات", "درس", "دروس", "في", "من", "على", "الى",
    "مع", "عن", "كامل", "بالتفصيل",
    "the", "of", "and", "for", "with", "part", "lecture", "lectures",
    "explained", "tutorial", "intro", "introduction",
)}


def _keys(text: str) -> set[str]:
    ws = re.findall(r"[\w\u0600-\u06FF]{3,}", text or "")
    return {_norm_w(w) for w in ws} - _STOP_W


# رموز تبقى معلّقة في نهاية العنوان بعد القصّ فتُحذف قبل إضافة النقاط
_TRAIL = " \t\u2026.,،؛;:|/\\-–—([{«\"'"


def _short(t: str, limit: int = 72) -> str:
    """قصّ عند حدّ الكلمة، بلا رمز معلّق ولا قوس مفتوح دون إغلاق."""
    t = re.sub(r"\s+", " ", t or "").strip()
    if len(t) <= limit:
        return t
    cut = t[:limit].rsplit(" ", 1)[0] or t[:limit]
    # قوس فُتح ثم قطعه القصّ: نحذفه وما بعده حتى لا يظهر «(أنواع الأجهزة…»
    for op, cl in (("(", ")"), ("[", "]"), ("«", "»")):
        if cut.count(op) > cut.count(cl):
            cut = cut[: cut.rfind(op)]
    cut = cut.rstrip(_TRAIL)
    return (cut or t[:limit].rstrip(_TRAIL) or t[:limit]) + "…"


def _is_school_level(item: dict) -> bool:
    txt = f"{item.get('title', '')} {item.get('channel', '')}"
    return bool(_SCHOOL_NOISE.search(txt))


def _off_topic(title: str, guard: set[str]) -> bool:
    """موضوع مجاور فعلاً، لا مادتنا؛ يُلغى الخصم إن كان جزءاً من المادة."""
    m = _NEAR_TOPIC.search(title or "")
    if not m:
        return False
    hit = {_norm_w(w) for w in re.findall(r"[\w\u0600-\u06FF]{3,}", m.group(0))}
    return not (hit & guard)


def _score(v: dict, tkeys: set[str], skeys: set[str], guard: set[str]) -> int:
    """كلما صغُرت القيمة تقدّم المقطع. الفرز مستقر فيبقى ترتيب يوتيوب عند التعادل."""
    title = v.get("title", "")
    s = 3 * len(tkeys & _keys(title))                            # تطابق موضوع الصفحة
    s += len(skeys & _keys(f"{title} {v.get('channel', '')}"))   # تطابق المادة
    s -= int(v.get("qi", 0))        # خصم خفيف: نتائج العبارة الثانية تنافس بعدالة
    if _BEGINNER.search(f"{title} {v.get('channel', '')}"):
        s -= 2                      # محتوى تعريفي للمبتدئين يتراجع أمام الجامعي
    if _off_topic(title, guard):
        s -= 3
    return -s


def _clean_topic(topic: str, subject: dict) -> str:
    """يقصّر الموضوع ليصلح عبارة بحث، ويعيد نصاً فارغاً إن كان عاماً."""
    t = re.sub(r"[\(\[].*?[\)\]]", " ", topic or "")   # المصطلح الإنكليزي يُستخرج منفصلاً
    t = _TOPIC_NOISE.sub("", t.strip())
    t = re.sub(r"[^\w\s\u0600-\u06FF+#.\-]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    words = [w for w in t.split(" ") if w]
    if len(words) > 6:
        t = " ".join(words[:6])
    t = t[:60].strip()

    tk = _keys(t)
    if not tk:
        return ""
    # الموضوع مطابق لاسم المادة ⇒ عبارة بحث عامة لا تصلح
    if tk <= _keys(f"{subject.get('name', '')} {subject.get('en', '')}"):
        return ""
    return t


def _split_topic(topic: str, subject: dict) -> tuple[str, str]:
    """يفصل «أجزاء العتاد (Computer Hardware)» إلى جزء عربي ومصطلح إنكليزي."""
    m = _PAREN_EN.search(topic or "")
    en_term = re.sub(r"\s+", " ", m.group(1)).strip() if m else ""
    return _clean_topic(topic, subject), en_term


def _mix(*parts: str) -> str:
    """يدمج أجزاء عبارة البحث ويُسقط أي جزء تكرّرت كل كلماته سلفاً."""
    out: list[str] = []
    seen: set[str] = set()
    for p in parts:
        p = (p or "").strip()
        if not p:
            continue
        pk = _keys(p)
        if pk and pk <= seen:
            continue
        seen |= pk
        out.append(p)
    return " ".join(out)


def _cap_query(q: str, max_words: int = 10, limit: int = 75) -> str:
    """يقصّ عبارة البحث عند حدّ الكلمة لا في منتصفها."""
    q = re.sub(r"\s+", " ", q or "").strip()
    w = q.split(" ")
    if len(w) > max_words:
        q = " ".join(w[:max_words])
    if len(q) <= limit:
        return q
    cut = q[:limit].rfind(" ")
    return (q[:cut] if cut > 20 else q[:limit]).strip()


def yt_queries(topic: str, subject: dict) -> list[str]:
    """عبارات بحث مرتبطة بموضوع الصفحة واسم المادة، بلا تكرار للكلمات."""
    t, en_term = _split_topic(topic, subject)
    en = (subject.get("en") or "").strip()
    ar = re.sub(r"\s*[\(\[].*?[\)\]]", "", subject.get("name", "")).strip()
    # المواد التقنية: المحتوى الإنكليزي الجامعي أوفر وأدق
    tech = bool(subject.get("is_code")) or subject.get("lang") == "en"

    if not t:                                   # موضوع عام أو مطابق لاسم المادة
        raw = [_mix("شرح", ar, "محاضرة جامعة"),
               f"{en} lecture" if en else _mix("شرح", ar)]
    elif has_arabic(t):
        raw = [
            # «جامعية» تُبعد نتائج المبتدئين والمناهج المدرسية عن العبارة الأولى
            _mix(f"شرح {t}", ar, "محاضرة جامعية"),
            _mix(en_term or t, en, "lecture") if (en_term or en) else f"شرح {t} جامعة",
            f"شرح {t} جامعة",
        ]
    elif tech:
        raw = [
            _mix(t, en, "lecture"),
            f"{t} explained",
            _mix(f"شرح {t}", ar),
        ]
    else:
        raw = [
            _mix(f"شرح {t}", ar),
            _mix(t, en) or f"{t} محاضرة",
            f"{t} explained",
        ]

    out, seen = [], set()
    for q in raw:
        q = _cap_query(q)
        if q and q.lower() not in seen:
            seen.add(q.lower())
            out.append(q)
    return out[:3]


async def yt_search(query: str, limit: int = YT_FETCH) -> list[dict]:
    """بحث واحد عبر YouTube Data API. يكلّف 100 وحدة مهما كان عدد النتائج."""
    if not YOUTUBE_API_KEY:
        return []
    params = {
        "part": "snippet", "q": query, "type": "video",
        "maxResults": limit, "safeSearch": "strict",
        "relevanceLanguage": "ar" if has_arabic(query) else "en",
        "key": YOUTUBE_API_KEY,
    }
    try:
        async with httpx.AsyncClient(timeout=YT_TIMEOUT) as client:
            r = await client.get(YT_API, params=params)
        if r.status_code != 200:
            logger.warning("youtube api %s: %s", r.status_code, r.text[:200])
            return []
        items = (r.json() or {}).get("items", [])
    except Exception as e:
        logger.warning("youtube search failed: %s", e)
        return []

    out = []
    for it in items:
        vid = (it.get("id") or {}).get("videoId")
        sn = it.get("snippet") or {}
        if not vid:
            continue
        # يوتيوب يعيد العنوان مرمّزاً (&amp; &quot; &#39;) وسنطبّق escape عند العرض،
        # لذا يُفكّ الترميز هنا مرة واحدة فقط لتفادي &amp;amp;
        out.append({
            "id": vid,
            "title": _h.unescape(sn.get("title") or "").strip(),
            "channel": _h.unescape(sn.get("channelTitle") or "").strip(),
            "url": f"https://www.youtube.com/watch?v={vid}",
        })
    return out


async def yt_search_many(queries: list[str], limit: int = YT_FETCH) -> list[dict]:
    """يجمع نتائج كل العبارات المسموحة مع تسجيل رقم العبارة.

    لا توقّف مبكراً: العبارة العربية تُشبع العدد دائماً فتمنع تنفيذ العبارة
    الإنكليزية، وهي غالباً الأدقّ في المواد التقنية.
    """
    out: list[dict] = []
    seen: set[str] = set()
    for qi, q in enumerate(queries[:YT_MAX_QUERIES]):
        key = f"{q.lower()}|{limit}"
        raw = _YT_CACHE.get(key)
        if raw is None:
            raw = await yt_search(q, limit)
            if raw:                       # لا نخزّن الفشل المؤقّت للشبكة
                if len(_YT_CACHE) >= _YT_CACHE_MAX:
                    _YT_CACHE.clear()
                _YT_CACHE[key] = raw
        for v in raw:
            if v["id"] in seen:
                continue
            seen.add(v["id"])
            item = dict(v)                # نسخة مستقلة: لا نكتب qi داخل الذاكرة المؤقتة
            item["qi"] = qi
            out.append(item)
    return out


async def build_youtube_block(topic: str, subject: dict) -> str:
    """كتلة يوتيوب بصيغة HTML: مقاطع حقيقية عند توفّر المفتاح، وإلا روابط بحث."""
    queries = yt_queries(topic, subject)
    if not queries:
        return ""

    videos: list[dict] = []
    if YOUTUBE_API_KEY:
        raw = await yt_search_many(queries)
        if raw:
            tkeys = _keys(topic)
            skeys = _keys(f"{subject.get('name', '')} {subject.get('en', '')}")
            guard = tkeys | skeys
            pool = [v for v in raw if not _is_school_level(v)] or raw
            pool.sort(key=lambda v: _score(v, tkeys, skeys, guard))
            videos = pool[:YT_RESULTS]

    if videos:
        lines = ["🎬 <b>مقاطع يوتيوب مقترحة</b>", ""]
        for v in videos:
            title = escape(_short(v["title"])) or "مقطع"
            lines.append(f"▫️ <a href=\"{v['url']}\">{title}</a>")
            if v["channel"]:
                lines.append(f"     <i>{escape(_short(v['channel'], 34))}</i>")
        more = YT_SEARCH_URL + quote_plus(queries[0])
        lines += ["", f"🔎 <a href=\"{more}\">مزيد من النتائج على يوتيوب</a>"]
    else:
        lines = ["🔎 <b>ابحث في يوتيوب عبر إحدى هذه العبارات</b>", ""]
        for q in queries:
            url = YT_SEARCH_URL + quote_plus(q)
            lines.append(f"▫️ <a href=\"{url}\">{escape(q)}</a>")

    # التنبيه العام انتقل إلى الرسالة الأخيرة (DISCLAIMER) فلا يُكرَّر هنا
    return "\n".join(lines)


async def send_plain(bot, chat_id: int, text: str, **kw) -> None:
    """إرسال نص بدون معاينة روابط، مع توافق مع إصدارات المكتبة."""
    try:
        if NO_PREVIEW is not None:
            await bot.send_message(chat_id, text, link_preview_options=NO_PREVIEW, **kw)
        else:
            await bot.send_message(chat_id, text, disable_web_page_preview=True, **kw)
    except TypeError:
        await bot.send_message(chat_id, text, **kw)


# ══════════════════════════════ معالجة الصور ══════════════════════════════

def compress_image(raw: bytes) -> bytes:
    if not PIL_OK:
        return raw
    try:
        im = Image.open(io.BytesIO(raw))
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        w, h = im.size
        if max(w, h) > IMG_MAX_SIDE:
            scale = IMG_MAX_SIDE / float(max(w, h))
            im = im.resize((max(1, int(w * scale)), max(1, int(h * scale))), Image.LANCZOS)
        out = io.BytesIO()
        q = IMG_QUALITY
        while True:
            out.seek(0)
            out.truncate(0)
            im.save(out, format="JPEG", quality=q, optimize=True)
            if out.tell() <= IMG_MAX_BYTES or q <= 45:
                break
            q -= 10
        return out.getvalue()
    except Exception as e:
        logger.warning("compress_image failed: %s", e)
        return raw


async def fetch_images(bot, file_ids: list[str]) -> list[str]:
    out: list[str] = []
    for fid in file_ids[:MAX_IMAGES]:
        f = await bot.get_file(fid)
        raw = bytes(await f.download_as_bytearray())
        small = await asyncio.to_thread(compress_image, raw)
        out.append(base64.b64encode(small).decode("ascii"))
    return out


# ══════════════════════════════ تنظيف النصوص ══════════════════════════════

_FENCE_RE = re.compile(r"(```.*?```)", re.DOTALL)
# سطر سياج كامل، مع وسم اللغة اختيارياً؛ يُستخدم عند بناء الملفات
_FENCE_LINE = re.compile(r"^\s*```+\s*[A-Za-z0-9+#._\-]*\s*$")


def _strip_reasoning(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<reasoning>.*?</reasoning>", "", text, flags=re.DOTALL | re.IGNORECASE)
    return text.strip()


def _clean_prose_segment(seg: str) -> str:
    seg = re.sub(r"\$\$(.*?)\$\$", r"\1", seg, flags=re.DOTALL)
    seg = re.sub(r"\$([^$\n]+)\$", r"\1", seg)
    seg = re.sub(r"\\?O\s*\(\s*n\s*\^\s*\{?2\}?\s*\)", "O(n²)", seg)
    seg = re.sub(r"\^\{?2\}?", "²", seg)
    seg = re.sub(r"\^\{?3\}?", "³", seg)
    seg = seg.replace("**", "").replace("__", "")
    seg = re.sub(r"^\s*#{1,6}\s+", "", seg, flags=re.MULTILINE)
    seg = re.sub(r"^\s{0,8}[\*\-]\s+", "• ", seg, flags=re.MULTILINE)
    return seg


def clean_code_output(text: str) -> str:
    """ينظّف الشرح ويترك ما بين ``` كما هو تماماً."""
    if not text:
        return ""
    text = _strip_reasoning(text)
    parts = _FENCE_RE.split(text)
    text = "".join(p if p.startswith("```") else _clean_prose_segment(p) for p in parts)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


_LATEX_MAP_AR = {
    r"\times": "×", r"\div": "÷", r"\cdot": "·", r"\pm": "±",
    r"\leq": "≤", r"\geq": "≥", r"\neq": "≠", r"\approx": "≈",
    r"\infty": "∞", r"\sum": "مجموع", r"\int": "تكامل",
    r"\alpha": "α", r"\beta": "β", r"\theta": "θ", r"\pi": "π",
    r"\mu": "μ", r"\sigma": "σ", r"\lambda": "λ", r"\Delta": "Δ",
    r"\left": "", r"\right": "", r"\,": " ", r"\;": " ", r"\!": "",
}

_LATEX_MAP_EN = dict(_LATEX_MAP_AR)
_LATEX_MAP_EN.update({r"\sum": "sum", r"\int": "integral"})


def clean_text_strictly(text: str, lang: str = "ar") -> str:
    """يحوّل LaTeX إلى رموز قابلة للعرض في تلغرام والوورد."""
    if not text:
        return ""
    text = _strip_reasoning(text)
    text = re.sub(r"\$\$(.*?)\$\$", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\$([^$\n]+)\$", r"\1", text)
    text = re.sub(r"\\\[(.*?)\\\]", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\\\((.*?)\\\)", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\\d?frac\s*\{([^{}]+)\}\s*\{([^{}]+)\}", r"(\1)/(\2)", text)
    text = re.sub(r"\\sqrt\s*\{([^{}]+)\}", r"√(\1)", text)
    text = re.sub(r"\\text\s*\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\math(?:bf|rm|it|cal)\s*\{([^{}]*)\}", r"\1", text)
    mapping = _LATEX_MAP_AR if lang == "ar" else _LATEX_MAP_EN
    for k, v in mapping.items():
        text = text.replace(k, v)
    text = re.sub(r"\^\{?2\}?", "²", text)
    text = re.sub(r"\^\{?3\}?", "³", text)
    text = re.sub(r"_\{([^{}]+)\}", r"_\1", text)
    text = re.sub(r"\\[a-zA-Z]+", "", text)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"^\s*#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s{0,8}[\*\-]\s+", "• ", text, flags=re.MULTILINE)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def clean_report_text(text: str, lang: str = "ar") -> str:
    """مثل clean_text_strictly لكنه يحافظ على علامة العنوان ## للتنسيق.

    كتل ``` تُترك حرفياً، ويُشترط فراغ بعد علامات # حتى لا يُفهم توجيه
    #include على أنه عنوان markdown فيتحوّل إلى ◾ ويتعطّل الكود.
    """
    if not text:
        return ""
    text = _strip_reasoning(text)
    out: list[str] = []
    for seg in _FENCE_RE.split(text):
        if seg.startswith("```"):
            out.append(seg)
            continue
        lines = []
        for line in seg.split("\n"):
            m = re.match(r"^\s*#{1,6}[ \t]+(.+?)\s*$", line)
            if m:
                lines.append("## " + clean_text_strictly(m.group(1), lang))
            else:
                lines.append(clean_text_strictly(line, lang))
        out.append("\n".join(lines))
    return re.sub(r"\n{3,}", "\n\n", "".join(out)).strip()


_PAREN_LATIN_RE = re.compile(r"\s*[\(\[]\s*[A-Za-z][A-Za-z0-9 ,.&'\-/]*\s*[\)\]]")


def strip_latin_parentheticals(text: str) -> str:
    """يحذف الأقواس التي تحوي مصطلحات لاتينية فقط، للمواد العربية.

    كتل ``` مستثناة، فحذف أقواس الكود يفسده.
    """
    if not text:
        return ""
    out: list[str] = []
    for seg in _FENCE_RE.split(text):
        if seg.startswith("```"):
            out.append(seg)
            continue
        seg = _PAREN_LATIN_RE.sub("", seg)
        seg = re.sub(r"[ \t]{2,}", " ", seg)
        out.append(re.sub(r"\s+([،.:؛])", r"\1", seg))
    return "".join(out).strip()


def headings_for_telegram(text: str) -> str:
    """يحوّل '## عنوان' إلى صيغة مقروءة في تلغرام، ولا يمسّ كتل الكود."""
    return "".join(
        p if p.startswith("```")
        else re.sub(r"^\s*##\s+(.+?)\s*$", r"◾ \1", p, flags=re.MULTILINE)
        for p in _FENCE_RE.split(text or "")
    )


_TOPIC_RE = re.compile(
    r"^\s*(?:الموضوع|العنوان|Topic|Title)\s*[:：]\s*(.+)$", re.MULTILINE)

# عناوين الأقسام الثابتة: ليست موضوعاً، وكانت تُعاد كموضوع فتفسد بحث يوتيوب
_GENERIC_HEADS = {_norm_w(h) for h in (
    "نظرة عامة", "المصطلحات الأساسية", "الشرح المفصل", "مثال محلول",
    "أخطاء شائعة", "أسئلة للمراجعة الذاتية", "ما ورد في الصفحة",
    "مفتاح الإجابات",
    "Overview", "Key Terms", "Detailed Explanation", "Worked Example",
    "Common Mistakes", "Self-Check Questions", "Answer Key",
)}


def extract_topic(body: str, fallback: str = "") -> str:
    """يستخرج عنوان الموضوع من ناتج الشرح ليُستخدم في البحث وفي اسم الملف."""
    m = _TOPIC_RE.search(body or "")
    if m:
        return re.sub(r"\s+", " ", m.group(1)).strip()[:120]
    for line in (body or "").split("\n"):
        s = re.sub(r"^\s*##\s*", "", line).strip()
        if len(s) >= 4 and _norm_w(s) not in _GENERIC_HEADS:
            return s[:120]
    return (fallback or "شرح")[:120]


def split_for_telegram(text: str, limit: int = TELEGRAM_CHUNK) -> list[str]:
    text = text.strip()
    if len(text) <= limit:
        return [text] if text else []
    chunks, rest = [], text
    while len(rest) > limit:
        window = rest[:limit]
        cut = window.rfind("\n\n")
        if cut < limit // 3:
            cut = window.rfind("\n")
        if cut < limit // 3:
            cut = window.rfind(" ")
        if cut < limit // 3:
            cut = limit
        chunks.append(rest[:cut].strip())
        rest = rest[cut:].strip()
    if rest:
        chunks.append(rest)
    return [c for c in chunks if c]


def iter_doc_blocks(body: str):
    """يقسّم المتن إلى ('head'|'text'|'code') للملفات.

    كتل الكود تُجمع كما هي بإزاحتها، وسطور السياج ``` لا تُطبع كنص.
    """
    in_code = False
    buf: list[str] = []
    for raw in (body or "").split("\n"):
        if _FENCE_LINE.match(raw):
            if in_code:
                while buf and not buf[0].strip():
                    buf.pop(0)
                while buf and not buf[-1].strip():
                    buf.pop()
                if buf:
                    yield ("code", buf)
                buf = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            buf.append(raw.rstrip())
            continue
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            yield ("head", line[3:].strip())
        else:
            yield ("text", line)
    if in_code:
        while buf and not buf[0].strip():
            buf.pop(0)
        while buf and not buf[-1].strip():
            buf.pop()
        if buf:
            yield ("code", buf)


# ══════════════════════════════ الإحصائيات ══════════════════════════════

_stats_lock = asyncio.Lock()
_stats_dirty = False


def _today_str() -> str:
    return datetime.now(BAGHDAD_TZ).strftime("%Y-%m-%d")


def _now_str() -> str:
    return datetime.now(BAGHDAD_TZ).isoformat(timespec="seconds")


def _empty_stats() -> dict:
    return {
        "schema": 2,
        "users": {},            # "معرّف" -> {first_seen, last_seen, name, username, requests}
        "requests_count": 0,
        "services_usage": {},
        "subjects_usage": {},
        "files": {"word": 0, "pdf": 0},
        "daily": {},            # "YYYY-MM-DD" -> {requests, users[]}
        "saved_at": None,
    }


def _migrate(data: dict) -> dict:
    """يقبل الصيغة القديمة (قائمة معرّفات) والجديدة ويعيد بنية موحّدة."""
    base = _empty_stats()
    if not isinstance(data, dict):
        return base

    users = data.get("users")
    if isinstance(users, list):
        for uid in users:
            base["users"][str(uid)] = {
                "first_seen": "", "last_seen": "", "name": "",
                "username": "", "requests": 0,
            }
    elif isinstance(users, dict):
        for uid, rec in users.items():
            if isinstance(rec, dict):
                base["users"][str(uid)] = {
                    "first_seen": rec.get("first_seen", "") or "",
                    "last_seen": rec.get("last_seen", "") or "",
                    "name": rec.get("name", "") or "",
                    "username": rec.get("username", "") or "",
                    "requests": int(rec.get("requests", 0) or 0),
                }

    try:
        base["requests_count"] = int(data.get("requests_count", 0) or 0)
    except Exception:
        pass

    for key in ("services_usage", "subjects_usage", "files"):
        src = data.get(key)
        if isinstance(src, dict):
            for k, v in src.items():
                try:
                    base[key][str(k)] = int(v or 0)
                except Exception:
                    continue

    daily = data.get("daily")
    if isinstance(daily, dict):
        for day, info in daily.items():
            if isinstance(info, dict):
                base["daily"][str(day)] = {
                    "requests": int(info.get("requests", 0) or 0),
                    "users": [str(u) for u in info.get("users", []) or []],
                }

    base["saved_at"] = data.get("saved_at")
    return base


def _prune_daily(data: dict) -> None:
    if len(data["daily"]) > KEEP_DAYS:
        for day in sorted(data["daily"])[:-KEEP_DAYS]:
            data["daily"].pop(day, None)


def _read_stats_file() -> dict:
    try:
        with open(STATS_FILE, "r", encoding="utf-8") as f:
            return _migrate(json.load(f))
    except FileNotFoundError:
        return _empty_stats()
    except Exception as e:
        logger.warning("stats read failed: %s", e)
        return _empty_stats()


def _write_stats_file(data: dict) -> None:
    try:
        os.makedirs(os.path.dirname(STATS_FILE) or ".", exist_ok=True)
        data["saved_at"] = _now_str()
        tmp = STATS_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        os.replace(tmp, STATS_FILE)
    except Exception as e:
        logger.warning("stats write failed: %s", e)


def _merge_stats(base: dict, incoming: dict) -> dict:
    """دمج آمن: يأخذ الأكبر لكل عدّاد واتحاد المستخدمين، فلا تتضخّم الأرقام بالتكرار."""
    inc = _migrate(incoming)

    for uid, rec in inc["users"].items():
        cur = base["users"].get(uid)
        if not cur:
            base["users"][uid] = rec
            continue
        cur["requests"] = max(cur.get("requests", 0), rec["requests"])
        if rec["first_seen"] and (not cur.get("first_seen")
                                  or rec["first_seen"] < cur["first_seen"]):
            cur["first_seen"] = rec["first_seen"]
        if rec["last_seen"] > cur.get("last_seen", ""):
            cur["last_seen"] = rec["last_seen"]
            cur["name"] = rec["name"] or cur.get("name", "")
            cur["username"] = rec["username"] or cur.get("username", "")

    for key in ("services_usage", "subjects_usage", "files"):
        for k, v in inc[key].items():
            base[key][k] = max(base[key].get(k, 0), v)

    for day, info in inc["daily"].items():
        cur = base["daily"].setdefault(day, {"requests": 0, "users": []})
        cur["requests"] = max(cur["requests"], info["requests"])
        cur["users"] = sorted(set(cur["users"]) | set(info["users"]))

    base["requests_count"] = max(base["requests_count"], inc["requests_count"])
    _prune_daily(base)
    return base


async def record_activity(user_id: int, service: str | None = None,
                          name: str = "", username: str = "",
                          subject_code: str | None = None) -> None:
    global _stats_dirty
    uid = str(user_id)
    async with _stats_lock:
        data = await asyncio.to_thread(_read_stats_file)
        rec = data["users"].get(uid)
        if rec is None:
            rec = {"first_seen": _now_str(), "last_seen": "", "name": "",
                   "username": "", "requests": 0}
            data["users"][uid] = rec
        rec["last_seen"] = _now_str()
        if name:
            rec["name"] = str(name)[:64]
        if username:
            rec["username"] = str(username)[:32]

        day = data["daily"].setdefault(_today_str(), {"requests": 0, "users": []})
        if uid not in day["users"]:
            day["users"].append(uid)

        if service:
            rec["requests"] = rec.get("requests", 0) + 1
            data["requests_count"] += 1
            data["services_usage"][service] = data["services_usage"].get(service, 0) + 1
            day["requests"] += 1
            if subject_code:
                data["subjects_usage"][subject_code] = \
                    data["subjects_usage"].get(subject_code, 0) + 1

        _prune_daily(data)
        await asyncio.to_thread(_write_stats_file, data)
        _stats_dirty = True


async def record_file(kind: str) -> None:
    """kind = 'word' أو 'pdf'."""
    global _stats_dirty
    async with _stats_lock:
        data = await asyncio.to_thread(_read_stats_file)
        data["files"][kind] = data["files"].get(kind, 0) + 1
        await asyncio.to_thread(_write_stats_file, data)
        _stats_dirty = True


async def remove_users(ids: list[int]) -> None:
    if not ids:
        return
    async with _stats_lock:
        data = await asyncio.to_thread(_read_stats_file)
        for uid in ids:
            data["users"].pop(str(uid), None)
        await asyncio.to_thread(_write_stats_file, data)


async def get_stats() -> dict:
    async with _stats_lock:
        return await asyncio.to_thread(_read_stats_file)


def build_stats_text(data: dict) -> str:
    users = data["users"]
    daily = data["daily"]
    today = _today_str()
    today_d = daily.get(today, {"requests": 0, "users": []})

    last7 = sorted(daily)[-7:]
    active7 = set()
    for d in last7:
        active7.update(daily[d].get("users", []))
    new_today = sum(1 for r in users.values()
                    if str(r.get("first_seen", ""))[:10] == today)

    lines = [
        "📊 إحصائيات المنصة الأكاديمية", "",
        f"الطلبة الكلي: {len(users)}",
        f"جديد اليوم: {new_today}",
        f"نشط اليوم: {len(today_d.get('users', []))}",
        f"نشط آخر سبعة أيام: {len(active7)}",
        f"مجموع العمليات: {data['requests_count']}",
        f"عمليات اليوم: {today_d.get('requests', 0)}", "",
        "الخدمات:",
    ]
    svc = sorted(data["services_usage"].items(), key=lambda x: -x[1])
    lines += [f"• {SERVICE_LABELS.get(k, k)}: {v}" for k, v in svc] or ["• لا يوجد"]

    subs = sorted(data["subjects_usage"].items(), key=lambda x: -x[1])[:8]
    if subs:
        lines += ["", "أكثر المواد طلباً:"]
        for code, count in subs:
            lines.append(f"• {SUBJECTS.get(code, {}).get('name', code)}: {count}")

    f = data["files"]
    lines += ["", f"الملفات المحمّلة: Word {f.get('word', 0)} — PDF {f.get('pdf', 0)}"]

    if last7:
        lines += ["", "آخر الأيام:"]
        for d in reversed(last7):
            info = daily[d]
            lines.append(f"• {d}: {info.get('requests', 0)} عملية / "
                         f"{len(info.get('users', []))} طالب")

    lines += ["", f"آخر حفظ: {data.get('saved_at') or 'غير معروف'}"]
    return "\n".join(lines)


# ───── النسخ الاحتياطي إلى القناة والاستعادة منها ─────

async def do_backup(bot, reason: str = "دوري") -> bool:
    if not BACKUP_CHAT_ID:
        return False
    try:
        data = await get_stats()
        blob = json.dumps(data, ensure_ascii=False, indent=1).encode("utf-8")
        caption = (f"نسخة احتياطية ({reason})\n{_now_str()}\n"
                   f"الطلبة: {len(data['users'])} | العمليات: {data['requests_count']}")
        msg = await bot.send_document(
            chat_id=BACKUP_CHAT_ID, document=io.BytesIO(blob),
            filename="bot_stats.json", caption=caption, disable_notification=True)
        try:
            await bot.unpin_all_chat_messages(chat_id=BACKUP_CHAT_ID)
        except Exception:
            pass
        try:
            await bot.pin_chat_message(chat_id=BACKUP_CHAT_ID,
                                       message_id=msg.message_id,
                                       disable_notification=True)
        except Exception as e:
            logger.warning("pin failed: %s", e)
        logger.info("backup sent (%s)", reason)
        return True
    except Exception as e:
        logger.warning("backup failed: %s", e)
        return False


async def restore_stats(bot) -> dict | None:
    """يقرأ الرسالة المثبّتة في قناة النسخ ويدمجها مع الملف المحلي."""
    if not BACKUP_CHAT_ID:
        return None
    try:
        chat = await bot.get_chat(BACKUP_CHAT_ID)
        pinned = getattr(chat, "pinned_message", None)
        if pinned is None or pinned.document is None:
            logger.info("no pinned backup found")
            return None
        f = await bot.get_file(pinned.document.file_id)
        raw = bytes(await f.download_as_bytearray())
        incoming = json.loads(raw.decode("utf-8"))
        async with _stats_lock:
            data = await asyncio.to_thread(_read_stats_file)
            data = _merge_stats(data, incoming)
            await asyncio.to_thread(_write_stats_file, data)
        logger.info("stats restored: %s users", len(data["users"]))
        return data
    except Exception as e:
        logger.warning("restore failed: %s", e)
        return None


async def backup_job(context: ContextTypes.DEFAULT_TYPE) -> None:
    global _stats_dirty
    if not _stats_dirty:
        return
    if await do_backup(context.bot, "دوري"):
        _stats_dirty = False


# ══════════════════════════════ ملف Word ══════════════════════════════

def _set_run_style(run, lang: str, size: int = 13, bold: bool = False,
                   color: tuple[int, int, int] | None = None,
                   font_name: str | None = None):
    name = font_name or (ARABIC_FONT if lang == "ar" else LATIN_FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)
    if lang == "ar":
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rPr.append(rtl)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(size * 2))
        rPr.append(szcs)


def _set_para_dir(paragraph, lang: str):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1" if lang == "ar" else "0")
    pPr.append(bidi)


def _add_par(doc, text: str, lang: str, size: int = 13, bold: bool = False,
             align=None, space_after: int = 8, color=None,
             font_name: str | None = None, line_spacing: float = 1.5):
    p = doc.add_paragraph()
    _set_para_dir(p, lang)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    run = p.add_run(text)
    _set_run_style(run, lang, size=size, bold=bold, color=color, font_name=font_name)
    return p


def _add_code_block(doc, lines: list[str]) -> None:
    """كتلة كود: خط ثابت العرض، يسار، سطر واحد، والإزاحة كما هي."""
    for ln in lines:
        _add_par(doc, ln if ln.strip() else "\u00a0", "en", size=10,
                 align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0,
                 font_name=MONO_FONT, line_spacing=1.0)


def create_academic_docx(title: str, body: str, subject_name: str, lang: str,
                         with_cover: bool = False, student: str = "") -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1.1)
    sec.right_margin = Inches(1.1)

    center = WD_ALIGN_PARAGRAPH.CENTER
    head_align = WD_ALIGN_PARAGRAPH.RIGHT if lang == "ar" else WD_ALIGN_PARAGRAPH.LEFT

    if with_cover:
        doc.add_paragraph()
        _add_par(doc, UNIVERSITY, "ar", size=20, bold=True, align=center, space_after=6)
        _add_par(doc, COLLEGE, "ar", size=15, align=center, space_after=4)
        _add_par(doc, DEPARTMENT, "ar", size=14, align=center, space_after=40)
        _add_par(doc, title, lang, size=24, bold=True, align=center,
                 space_after=30, color=(0x1F, 0x3B, 0x73))
        _add_par(doc, subject_name, "ar", size=14, align=center, space_after=6)
        _add_par(doc, STAGE, "ar", size=13, align=center, space_after=6)
        if student:
            _add_par(doc, student, "ar", size=13, align=center, space_after=6)
        _add_par(doc, academic_year(), "ar", size=13, align=center, space_after=0)
        doc.add_page_break()
    else:
        _add_par(doc, title, lang, size=19, bold=True, align=center,
                 space_after=18, color=(0x1F, 0x3B, 0x73))

    for kind, item in iter_doc_blocks(body):
        if kind == "code":
            _add_code_block(doc, item)
        elif kind == "head":
            _add_par(doc, item, lang, size=15, bold=True,
                     align=head_align, space_after=6, color=(0x1F, 0x3B, 0x73))
        else:
            _add_par(doc, item, lang, size=13)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ══════════════════════════════ ملف PDF ══════════════════════════════
# ملاحظة: الملف قد يخلط العربية واللاتينية (عنوان عربي مع متن إنكليزي)،
# لذلك يُسجَّل خط يونيكودي واحد لكل الأنماط، ويُحدَّد اتجاه كل فقرة وحدها.

PDF_FAMILY = "doc"
PDF_MONO = "mono"
_PDF_MAX_WORD = 42

_AR_CHARS = re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF"
                       r"\uFB50-\uFDFF\uFE70-\uFEFF]")


def has_arabic(text: str) -> bool:
    return bool(_AR_CHARS.search(text or ""))


# رموز شائعة قد لا يحتويها الخط، تُستبدل ببدائل آمنة
_PDF_MAP = {
    "\u2014": "-", "\u2013": "-", "\u2012": "-", "\u2212": "-", "\u2015": "-",
    "\u201c": '"', "\u201d": '"', "\u201e": '"',
    "\u2018": "'", "\u2019": "'", "\u2032": "'", "\u2033": '"',
    "\u2026": "...", "\u2022": "-", "\u25cf": "-", "\u25aa": "-", "\u25a0": "-",
    "\u2192": "->", "\u2190": "<-", "\u21d2": "=>", "\u21d4": "<=>",
    "\u221a": "sqrt", "\u221e": "inf", "\u2248": "~", "\u2260": "!=",
    "\u2264": "<=", "\u2265": ">=", "\u2211": "sum", "\u222b": "integral",
    "\u2208": "in", "\u2229": "and", "\u222a": "or", "\u00ac": "not",
    "\u03b1": "alpha", "\u03b2": "beta", "\u03b8": "theta", "\u03c0": "pi",
    "\u03bc": "mu", "\u03c3": "sigma", "\u03bb": "lambda", "\u0394": "Delta",
    "\u03a3": "Sigma", "\u03b3": "gamma", "\u03c1": "rho", "\u03c6": "phi",
    "\u00a0": " ", "\u202f": " ", "\u2009": " ", "\u200a": " ",
    "\u200b": "", "\u200c": "", "\u200d": "", "\u200e": "", "\u200f": "",
    "\u202a": "", "\u202b": "", "\u202c": "", "\ufeff": "",
    "\u2028": " ", "\u2029": " ", "\t": "    ",
}


def sanitize_for_pdf(text: str, keep_spaces: bool = False) -> str:
    """يستبدل الرموز غير المدعومة، يحذف الإيموجي، ويقطّع الكلمات الطويلة جداً.

    keep_spaces=True لكتل الكود: لا تُدمج الفراغات المتتالية فتبقى الإزاحة.
    """
    if not text:
        return ""
    for bad, good in _PDF_MAP.items():
        text = text.replace(bad, good)

    keep = []
    for ch in text:
        o = ord(ch)
        if ch in "\n\r":
            keep.append(ch)
        elif o < 0x20:
            keep.append(" ")
        elif 0x2300 <= o <= 0x2BFF or 0xE000 <= o <= 0xF8FF or o >= 0x1F000:
            continue                      # رسوم ورموز وإيموجي
        elif 0xFE00 <= o <= 0xFE0F:
            continue                      # محدّدات شكل الإيموجي
        else:
            keep.append(ch)
    text = "".join(keep)

    parts = []
    for word in text.split(" "):
        while len(word) > _PDF_MAX_WORD:
            parts.append(word[:_PDF_MAX_WORD])
            word = word[_PDF_MAX_WORD:]
        parts.append(word)
    text = " ".join(parts)
    if keep_spaces:
        return text
    return re.sub(r"[ ]{3,}", "  ", text)


def _hard_clean(text: str) -> str:
    """آخر خط دفاع: لاتيني أساسي وعربي فقط."""
    keep = []
    for ch in text:
        o = ord(ch)
        if ch in "\n\r" or 0x20 <= o <= 0x7E or 0x0600 <= o <= 0x06FF \
                or 0xFB50 <= o <= 0xFEFF:
            keep.append(ch)
    return "".join(keep) or " "


def _pdf_ready(lang: str = "ar") -> tuple[bool, str]:
    if not FPDF_OK:
        return False, "مكتبة fpdf2 غير مثبّتة على الخادم."
    if not os.path.exists(AR_TTF):
        return False, f"ملف الخط غير موجود في المسار {AR_TTF}."
    return True, ""


class _ReportPDF(FPDF):
    def __init__(self, lang: str):
        super().__init__(format="A4", unit="mm")
        self.lang = lang
        self.base_font = PDF_FAMILY
        self.set_margins(20, 20, 20)
        self.set_auto_page_break(True, margin=20)

        # خط واحد يونيكودي لكل الأنماط، فلا رجوع إلى helvetica مطلقاً
        bold_path = AR_TTF_BOLD if os.path.exists(AR_TTF_BOLD) else AR_TTF
        self.add_font(PDF_FAMILY, "", AR_TTF)
        self.add_font(PDF_FAMILY, "B", bold_path)
        self.add_font(PDF_FAMILY, "I", AR_TTF)
        self.add_font(PDF_FAMILY, "BI", bold_path)

        # خط ثابت العرض للكود إن وُجد، وإلا يُستخدم الخط الأساسي
        self.mono_ok = False
        if os.path.exists(MONO_TTF):
            try:
                self.add_font(PDF_MONO, "", MONO_TTF)
                self.add_font(PDF_MONO, "B", MONO_TTF)
                self.mono_ok = True
            except Exception as e:
                logger.warning("mono font not loaded: %s", e)

        self.shaped = False
        self._dir = None
        try:
            self.set_text_shaping(True, direction="ltr", script="latn", language="eng")
            self.shaped = True
            self._dir = "ltr"
        except Exception as e:
            logger.warning("text shaping unavailable: %s", e)

        self.use(12)

    def use(self, size: int, bold: bool = False):
        self.set_font(self.base_font, "B" if bold else "", size)

    def use_mono(self, size: int):
        self.set_font(PDF_MONO if self.mono_ok else self.base_font, "", size)

    def _set_dir(self, rtl: bool) -> None:
        if not self.shaped:
            return
        want = "rtl" if rtl else "ltr"
        if self._dir == want:
            return
        try:
            if rtl:
                self.set_text_shaping(True, direction="rtl",
                                      script="arab", language="ara")
            else:
                self.set_text_shaping(True, direction="ltr",
                                      script="latn", language="eng")
            self._dir = want
        except Exception as e:
            logger.warning("set direction failed: %s", e)
            self.shaped = False

    def _prep(self, text: str, rtl: bool) -> str:
        if not rtl or self.shaped:
            return text
        if BIDI_OK:
            try:
                return get_display(arabic_reshaper.reshape(text))
            except Exception:
                return text
        return text

    def _wrap(self, text: str, rtl: bool) -> list[str]:
        avail = self.w - self.l_margin - self.r_margin
        lines, cur = [], ""
        for word in text.split(" "):
            trial = word if not cur else cur + " " + word
            if self.get_string_width(self._prep(trial, rtl)) <= avail or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = word
        if cur:
            lines.append(cur)
        return lines or [""]

    def _render(self, text: str, rtl: bool, center: bool, lh: float,
                code: bool = False) -> None:
        self._set_dir(False if code else rtl)
        if code:
            self.multi_cell(0, lh, text, align="L", new_x="LMARGIN", new_y="NEXT")
            return
        if self.shaped:
            align = "C" if center else ("R" if rtl else "J")
            self.multi_cell(0, lh, text, align=align, new_x="LMARGIN", new_y="NEXT")
        else:
            align = "C" if center else ("R" if rtl else "L")
            for line in self._wrap(text, rtl):
                self.cell(0, lh, self._prep(line, rtl), align=align,
                          new_x="LMARGIN", new_y="NEXT")

    def par(self, text: str, size: int = 12, bold: bool = False,
            center: bool = False, gap: float = 3.0, rgb=None, code: bool = False):
        text = sanitize_for_pdf(text, keep_spaces=code)
        if not text.strip() and not code:
            return
        rtl = False if code else has_arabic(text)
        if code:
            self.use_mono(size)
        else:
            self.use(size, bold)
        self.set_text_color(*(rgb or (0, 0, 0)))
        lh = max(4.5, size * 0.62)
        try:
            self._render(text, rtl, center, lh, code=code)
        except Exception as e:
            logger.warning("pdf paragraph failed (%s) — retrying cleaned", e)
            try:
                self._render(_hard_clean(text), rtl, center, lh, code=code)
            except Exception as e2:
                logger.warning("pdf paragraph dropped: %s", e2)
        self.ln(gap)
        self.set_text_color(0, 0, 0)


def create_academic_pdf(title: str, body: str, subject_name: str, lang: str,
                        with_cover: bool = False, student: str = "") -> bytes:
    pdf = _ReportPDF(lang)
    pdf.add_page()

    if with_cover:
        pdf.ln(25)
        pdf.par(UNIVERSITY, size=22, bold=True, center=True, gap=3)
        pdf.par(COLLEGE, size=15, center=True, gap=2)
        pdf.par(DEPARTMENT, size=14, center=True, gap=35)
        pdf.par(title, size=24, bold=True, center=True, gap=25, rgb=(31, 59, 115))
        pdf.par(subject_name, size=14, center=True, gap=3)
        pdf.par(STAGE, size=13, center=True, gap=3)
        if student:
            pdf.par(student, size=13, center=True, gap=3)
        pdf.par(academic_year(), size=13, center=True, gap=0)
        pdf.add_page()
    else:
        pdf.par(title, size=19, bold=True, center=True, gap=8, rgb=(31, 59, 115))

    for kind, item in iter_doc_blocks(body):
        if kind == "code":
            pdf.ln(1)
            for ln in item:
                pdf.par(ln if ln.strip() else " ", size=9, gap=0, code=True)
            pdf.ln(3)
        elif kind == "head":
            pdf.ln(2)
            pdf.par(item, size=15, bold=True, gap=2, rgb=(31, 59, 115))
        else:
            pdf.par(item, size=12, gap=3)

    return bytes(pdf.output())


def safe_filename(text: str, fallback: str = "report") -> str:
    text = re.sub(r"[\\/:*?\"<>|\n\r\t]+", " ", text).strip()
    text = re.sub(r"\s{2,}", " ", text)[:60]
    return text or fallback


# ══════════════════════════════ التعليمات ══════════════════════════════

AR_ONLY_RULE = (
    "قاعدة إلزامية: لا تكتب أي كلمة إنكليزية ولا أي حرف لاتيني إطلاقاً. "
    "لا تضع مقابلاً إنكليزياً بين قوسين لأي مصطلح، ولا أسماء أعلام بالإنكليزية، "
    "ولا مصطلحات مترجمة داخل أقواس. استخدم المصطلح العربي وحده دائماً. "
    "ويُستثنى من هذا المنع الأرقام والرموز الرياضية ورموز المتغيّرات داخل "
    "المعادلات وحدها، فاكتبها كما هي دون ترجمة. "
)

# عامة بلا أمثلة مصطلحات: الأمثلة تتسرّب إلى الشرح فيذكر النموذج مفاهيم
# غير موجودة في الصفحة (Cache) أو أخطاءً لم يقع فيها الطالب («الحافلة»)
TERM_RULE = (
    "قاعدة المصطلحات: لا تترجم المصطلحات التقنية ترجمة حرفية قد تُفهم خطأ. "
    "اكتب المصطلح العربي المتعارف عليه أكاديمياً ثم الإنكليزي بين قوسين، "
    "وإن لم يوجد مقابل عربي دقيق فأبقِ المصطلح الإنكليزي كما هو. "
    "كلمة Lecture تعني «المحاضرة» لا «الدورة»، وChapter تعني «الفصل»."
)

# تُضاف في مسار النص للمواد التقنية فقط: أمثلتها عن العتاد، وكانت تُحقن
# في مواد مثل اللغة الإنكليزية والإحصاء فتُدخل مفاهيم غريبة عن الدرس
TERM_EXAMPLES = (
    " أمثلة على الصياغة المطلوبة: الناقل (Bus) لا «الحافلة»، "
    "الذاكرة المخبئية (Cache) لا «المخفية»، السجل (Register) لا «التسجيل»، "
    "المقاطعة (Interrupt)، النواة (Core)."
)

# نسخة آمنة لمسار الصورة: موسومة صراحة بأنها صياغة لا محتوى، فلا يتسرّب
# منها مفهوم إلى الشرح. أُضيفت بعد ظهور «المحيطات» و«الغلاف» في الشرح.
TERM_EXAMPLES_SAFE = (
    " أمثلة على الصياغة اللغوية فقط، وهي ليست محتوى مطلوباً، فلا تذكر أي مفهوم "
    "منها إن لم يرد في الصفحة: Peripherals = الوحدات الطرفية لا «المحيطات»، "
    "Bus = الناقل لا «الحافلة»، Case = الهيكل لا «الغلاف»، "
    "Register = السجل، Interrupt = المقاطعة، "
    "Volatile = متطايرة و Non-volatile = غير متطايرة، لا «متقطعة» ولا «متقلبة». "
    "والتزم بمقابل عربي واحد لكل مصطلح في الشرح كله ولا تنوّع صياغته."
)

# أُضيفت بعد ظهور أمثلة عربية في شرح مادة اللغة الإنكليزية: الطالب يُسأل
# عن المثال بلغة المادة، فالمثال العربي وحده لا ينفعه في الامتحان
EXAMPLE_LANG_RULE = (
    " قاعدة لغة الأمثلة: كل مثال أو جملة توضيحية أو نص للتحليل يُكتب أولاً بلغة "
    "دراسة المادة، ثم تُذكر ترجمته العربية في سطر تالٍ. ولا تستبدل المثال بمثال "
    "عربي ولا تكتفِ به، لأن الطالب سيُسأل عن المثال بلغة المادة في الامتحان. "
    "أما الشرح والتحليل والتعليق فتبقى بالعربية. "
)

# حدّ للمسار النصي: لا صفحة تُقيّد الشرح، فيُقيَّد بالموضوع المطلوب وحده
SCOPE_RULE = (
    " قاعدة النطاق: اقتصر على الموضوع المطلوب ولا تستطرد إلى مفاهيم أخرى من "
    "المادة أو من الفصول المجاورة. وإن كان ذكر مفهوم خارجي ضرورياً لإكمال المعنى "
    "فاكتفِ بجملة واحدة عنه وصرّح بأنه خارج نطاق الموضوع. "
    "وإن بنيت مثالاً أو تصنيفاً على استنتاج فصرّح بأنه استنتاج من إنشائك. "
    "وقبل الإخراج راجع سلامة العربية نحوياً، وخاصة مطابقة الصفة والحال لما "
    "يعودان إليه في الجنس والعدد. "
)

GROUND_RULE = (
    "قاعدة الالتزام بالمصدر: اشرح ما ورد في الصفحة فقط، ولا تُدخل مفاهيم أو مصطلحات "
    "غير مذكورة فيها ولو كانت من المادة نفسها أو من الفصول المجاورة لها. "
    "وإن كان ذكر مفهوم خارجي ضرورياً لإكمال المعنى فاكتفِ بجملة واحدة عنه وصرّح "
    "بأنه خارج نطاق هذه الصفحة. "
    "وانقل أسماء الأعلام وأسماء المواد وأرقام المحاضرات والفصول حرفياً كما وردت في "
    "الصورة دون تعديل حرف واحد، وإن لم تقرأها بثقة فلا تذكرها إطلاقاً. "
    "ولا تتوقّع أسئلة الامتحان ولا تقل إن مسألة ما تأتي في الامتحان أو أنها من أكثر "
    "ما يُسأل، فأنت لا تعرف امتحان هذه المادة ولا أسلوب تدريسها. "
    "وإن بنيت مثالاً أو تصنيفاً على استنتاج لم تنصّ عليه الصفحة فصرّح بأنه استنتاج "
    "ولا تنسبه إلى الصفحة. "
    "واكتفِ بتشبيه واحد في الشرح كله ولا تكتب فقرة مستقلة للتشبيه. "
    "وقبل الإخراج راجع سلامة العربية نحوياً، وخاصة مطابقة الصفة والحال لما يعودان "
    "إليه في الجنس والعدد."
)

# أُضيفت بعد أخطاء تحويل الوحدات (خطأ بمعامل 1024) وقلب التسلسل الهرمي للذاكرة
NUM_RULE = (
    "قاعدة الأرقام: إن أوردت مثالاً حسابياً فاكتب تحويل الوحدات صريحاً في كل خطوة "
    "(مثال: 1 GB/s = 1024 MB/s = 1048576 KB/s)، ولا تنتقل إلى خطوة تالية قبل إثبات "
    "التحويل، ثم تحقّق من أن رتبة النتيجة منطقية (نانوثانية، ميكروثانية، ميلي ثانية). "
    "ولا تختلق مواصفات عددية لأجهزة بعينها، واكتفِ بمراتب تقريبية إن لم تكن الأرقام "
    "مذكورة في المصدر. "
    "والتزم بترتيب التسلسل الهرمي للذاكرة من الأسرع إلى الأبطأ: السجلات، ثم الذاكرة "
    "المخبئية، ثم الذاكرة الرئيسية، ثم أقراص الحالة الصلبة، ثم الأقراص الميكانيكية، "
    "ولا تعكس هذا الترتيب في أي حال."
)

# أُضيفت بعد اختلاق عبارات وأبيات في مادة البلاغة ونسبتها إلى الشعر
QUOTE_RULE = (
    "قاعدة الشواهد: لا تختلق أبياتاً شعرية ولا أمثالاً ولا عبارات تراثية ولا آيات "
    "ولا أحاديث. ولا تستشهد بنصّ إلا إن كنت واثقاً من لفظه ونسبته معاً، وإلا فاصنع "
    "مثالاً من إنشائك وصرّح بأنه مثال توضيحي لا شاهد أدبي. "
    "وفي المواد غير التقنية (النحو، البلاغة، الأدب، الإدارة) اقتصر على القواعد "
    "المتّفق عليها، وتحقّق من صحة كل مثال إعرابياً وبلاغياً قبل إخراجه، ولا تصف "
    "تركيباً عربياً سليماً بأنه خطأ."
)

OCR_PREFIX = (
    "الصورة المرفقة تحتوي على سؤال أو مسألة أكاديمية. "
    "اقرأ كل ما فيها بدقة: النصوص والأرقام والرموز والمخططات والجداول. "
    "ابدأ ردك بقسم عنوانه «المعطيات كما قرأتها من الصورة» تُدرج فيه ما استخرجته، "
    "ثم قدّم الحل، ثم اختم بملاحظة تطلب من الطالب التأكد من مطابقة المعطيات لما في ورقته. "
    "إذا كان جزء غير واضح فاذكر ذلك صريحاً ولا تخترع أرقاماً."
)

OCR_CODE_EXTRA = (
    "تحذير مهم قبل القراءة: الصورة قد تكون مصوّرة من شاشة بواجهة عربية، "
    "فتظهر الأقواس والفواصل المنقوطة في مواضع مقلوبة بسبب اتجاه العرض من اليمين إلى اليسار. "
    "الكود البرمجي يُقرأ دائماً من اليسار إلى اليمين، وأي فاصلة منقوطة أو قوس يبدو في بداية السطر "
    "هو في الحقيقة في نهايته. لا تعتبر مواضع الترقيم المقلوبة أخطاءً من الطالب ولا تذكرها في قسم "
    "الأخطاء، واقتصر على الأخطاء المنطقية والدلالية الحقيقية."
)

OCR_EXPLAIN = (
    "الصور المرفقة صفحات من ملزمة أو كتاب منهجي أو سلايدات محاضرة جامعية. "
    "اقرأ كل ما فيها بدقة: النصوص والعناوين والمعادلات والجداول والمخططات، "
    "وإن كانت الصفحات متعددة فاعتبرها متسلسلة ومترابطة. "
    "مهمّتك شرح هذا المحتوى للطالب شرحاً مبسّطاً ومتدرّجاً لا مجرد إعادة كتابته. "
    "ابدأ ردك بسطر واحد على هذه الصورة: «الموضوع: ...» يصف المحتوى الفعلي "
    "لهذه الصفحة تحديداً في خمس كلمات أو أقل، مع المصطلح الإنكليزي بين قوسين. "
    "لا تكتب اسم المادة ولا عنوان المحاضرة العام، بل الفكرة المحدّدة التي تعالجها "
    "الصفحة. مثال: إذا كانت الصفحة عن تصنيف أجزاء العتاد وأجهزة الإدخال فاكتب "
    "«الموضوع: أجزاء العتاد وأجهزة الإدخال (Computer Hardware Parts)» "
    "لا «الموضوع: معمارية الحاسوب». هذا السطر يُستخدم للبحث عن مقاطع تعليمية، "
    "فكلما كان دقيقاً كانت المقاطع أقرب لدرسك. "
    "ثم اكتب قسماً بعنوان '## ما ورد في الصفحة' تُلخّص فيه العناوين "
    "والنقاط الأساسية كما هي، ثم أكمل بأقسام الشرح المطلوبة. "
    + TERM_RULE
    + TERM_EXAMPLES_SAFE +
    " إذا كان جزء من الصورة غير واضح فاذكر ذلك صريحاً ولا تخترع محتوى غير موجود، "
    "ولا تضف معلومات تخرج عن موضوع الصفحة إلا إذا كانت ضرورية لفهمها. "
    + NUM_RULE
    + GROUND_RULE + " "
    + QUOTE_RULE          # كانت غائبة عن مسار الصورة، فصفحة بلاغة مصوّرة بلا حماية
)


def explain_lang(subject: dict) -> str:
    """لغة الشرح: عربية افتراضياً حتى في المواد الإنكليزية، للفهم الأسرع."""
    return "ar" if EXPLAIN_IN_ARABIC else subject.get("lang", "ar")


def sys_report(subject: dict) -> str:
    if subject["lang"] == "en":
        return (
            "You are an academic writing assistant for a second-year Computer Science student. "
            f"Write a well-structured academic report for the course: {subject['en']}. "
            "Write the ENTIRE report in clear academic English. "
            "Structure: an introduction, several body sections, and a conclusion. "
            "Mark every section heading with '## ' at the start of its own line. "
            "Use plain paragraphs only: no markdown bold, no bullet symbols, no tables, "
            "no LaTeX and no dollar signs. Write formulas in plain notation such as O(n^2) "
            "or (a+b)/2. Do not include a cover page, student name, university name or date. "
            "Do not include source code blocks. Aim for 700-1100 words."
        )
    return (
        "أنت مساعد كتابة أكاديمية لطالب في المرحلة الثانية بعلوم الحاسوب. "
        f"اكتب تقريراً أكاديمياً متكاملاً في مادة: {subject['name']}. "
        "اكتب التقرير كاملاً بلغة عربية فصحى واضحة. "
        "البنية: مقدمة، ثم عدة محاور، ثم خاتمة. "
        "ضع علامة '## ' في بداية سطر كل عنوان فرعي. "
        "استخدم فقرات نصية فقط: بلا markdown، بلا نجوم، بلا جداول، بلا LaTeX وبلا علامات دولار. "
        + AR_ONLY_RULE +
        "لا تكتب صفحة غلاف ولا اسم جامعة ولا تاريخ. لا تُدرج أكواداً برمجية. "
        "استهدف 700 إلى 1100 كلمة."
    )


def sys_explain(subject: dict) -> str:
    """شرح تعليمي مبسّط لموضوع أو لمحتوى صفحة ملزمة."""
    lang = explain_lang(subject)

    if lang == "en":
        return (
            "You are a patient private tutor for a second-year Computer Science student "
            f"in the course: {subject['en']}. "
            "Explain the requested topic from scratch, assuming the student is confused and "
            "needs clarity, not a formal report. Build the idea step by step, from the simple "
            "to the more advanced, and use at most one concrete everyday analogy in the whole "
            "explanation. "
            "Use exactly these headings, each on its own line starting with '## ':\n"
            "## Overview\n## Key Terms\n## Detailed Explanation\n## Worked Example\n"
            "## Common Mistakes\n## Self-Check Questions\n"
            "In Key Terms give a one-line definition per term. In Worked Example show every "
            "step. In Self-Check Questions give four short questions with brief answers. "
            "Plain text only: no markdown emphasis, no tables, no LaTeX, no dollar signs; "
            "write math in plain notation such as O(n^2), sqrt(x), (a+b)/2. "
            "In any numerical example, show every unit conversion explicitly and sanity-check "
            "the order of magnitude of the result; never invent device specifications, and "
            "never claim that a storage device is faster than main memory. "
            "Never invent quotations, verses or proverbs. "
            "Never invent references, page numbers or book names. "
            "Stay within the requested topic and do not drift into neighbouring chapters."
        )

    base = (
        "أنت مدرّس خصوصي صبور لطالب في المرحلة الثانية بعلوم الحاسوب، "
        f"والمادة هي: {subject['name']}. "
        "اشرح الموضوع المطلوب من الصفر بأسلوب تعليمي مبسّط، لا بأسلوب تقرير رسمي. "
        "ابنِ الفكرة خطوة بخطوة من الأسهل إلى الأصعب، واستخدم تشبيهاً من الحياة اليومية "
        "حين يساعد ذلك على الفهم. "
        "واكتفِ بتشبيه واحد في الشرح كله، ولا تكتب فقرة مستقلة للتشبيه، ولا تستخدم "
        "التشبيه نفسه لمفهومين مختلفين أو متعاكسين حتى لا يلتبس الفرق بينهما على الطالب. "
        "ولا تتوقّع أسئلة الامتحان ولا تقل إن مسألة ما من أكثر ما يُسأل. "
        "استخدم هذه العناوين حرفياً، كل عنوان في سطر مستقل يبدأ بـ '## ':\n"
        "## نظرة عامة\n## المصطلحات الأساسية\n## الشرح المفصل\n## مثال محلول\n"
        "## أخطاء شائعة\n## أسئلة للمراجعة الذاتية\n"
        "في المصطلحات الأساسية اكتب تعريفاً في سطر واحد لكل مصطلح. "
        "في المثال المحلول اعرض كل خطوة بالتفصيل. "
        "في أسئلة المراجعة اكتب أربعة أسئلة قصيرة مع إجابة موجزة لكل سؤال، "
        "وتأكّد من صحة كل إجابة قبل إخراجها. "
        "نص عادي فقط: بلا markdown وبلا جداول وبلا LaTeX وبلا علامات دولار، "
        "واكتب الرياضيات برموز عادية مثل O(n^2) و √(x) و (أ+ب)/2. "
        "لا تختلق مراجع ولا أرقام صفحات ولا أسماء كتب. "
        + NUM_RULE + " "
        + QUOTE_RULE + " "
        + SCOPE_RULE
    )

    if subject.get("lang") == "en":
        tail = (
            "اكتب الشرح بالعربية الفصحى، لكن أبقِ كل مصطلح تقني بصيغته الإنكليزية الأصلية "
            "مع ذكر مقابله العربي بين قوسين عند أول ظهور فقط، مثل: Cache Memory (الذاكرة المخبئية). "
            "أسماء الأوامر والدوال والسجلات تبقى بالإنكليزية كما هي دائماً. "
        )
        tail += EXAMPLE_LANG_RULE + TERM_RULE
        if subject.get("is_code"):
            tail += TERM_EXAMPLES      # أمثلة العتاد للمواد التقنية وحدها
        return base + tail
    return base + AR_ONLY_RULE


def sys_math(subject: dict) -> str:
    if subject["lang"] == "en":
        return (
            f"You are a teaching assistant for the course {subject['en']} (second-year CS). "
            "Solve the problem step by step. Show the given data, the method, every intermediate "
            "computation, and the final answer clearly labelled. Verify the result at the end. "
            "Write in English. Never use LaTeX or dollar signs; write math in plain notation "
            "(x^2, sqrt(x), (a+b)/2, 3.16). Show every unit conversion explicitly and check that "
            "the order of magnitude of the result is plausible. If the problem is statistical, "
            "state explicitly whether you used the sample or the population formula."
        )
    return (
        f"أنت مدرّس مساعد لمادة {subject['name']} للمرحلة الثانية. "
        "حل المسألة خطوة بخطوة: اذكر المعطيات، ثم الطريقة، ثم كل خطوة حسابية، "
        "ثم النتيجة النهائية بوضوح، واختم بتحقق من صحتها. "
        "اكتب بالعربية ولا تستخدم LaTeX ولا علامات الدولار إطلاقاً؛ اكتب الرياضيات برموز عادية "
        "مثل س² و √(س) و (أ+ب)/2. "
        "واكتب كل تحويل للوحدات صريحاً وتحقّق من أن رتبة النتيجة منطقية. "
        "وإن كانت المسألة إحصائية فصرّح هل استخدمت صيغة العيّنة أم المجتمع. "
        "لا تكتب أي كلمة إنكليزية ولا حرفاً لاتينياً في الشرح، واستخدم المصطلحات العربية وحدها "
        "(الأرقام والرموز الرياضية مستثناة من هذا المنع)."
    )


def sys_code(subject: dict) -> str:
    return (
        f"You are a code reviewer for the course {subject['en']} (second-year CS students). "
        "Given the student's code: (1) state its purpose briefly, (2) list the real bugs one by "
        "one with the reason, (3) provide the full corrected code inside a single fenced block "
        "with the correct language tag, (4) state the time and space complexity. "
        "Write the explanation in Arabic but keep all identifiers, keywords and code in their "
        "original language. Never alter escape sequences such as \\n or \\t inside strings. "
        "Write complexity in plain form like O(n²) or O(1) with no dollar signs and no LaTeX. "
        "If the code has no real errors, say so plainly instead of inventing problems."
    )


def sys_quiz(subject: dict) -> str:
    if subject["lang"] == "en":
        return (
            f"You are an exam preparation assistant for the course {subject['en']}. "
            "Produce a question bank on the requested topic in English: five multiple-choice "
            "questions, three short-answer questions and two applied/analytical questions. "
            "Then add a separate answer key section marked with '## Answer Key'. "
            "Verify every answer before writing it, and never invent quotations or references. "
            "Use plain text, no markdown emphasis, no LaTeX, no dollar signs."
        )
    return (
        f"أنت مساعد تحضير للامتحانات في مادة {subject['name']}. "
        "أنشئ بنك أسئلة حول الموضوع المطلوب بالعربية: خمسة أسئلة اختيار من متعدد، "
        "وثلاثة أسئلة قصيرة، وسؤالان تطبيقيان تحليليان. "
        "ثم أضف قسماً منفصلاً يبدأ بـ '## مفتاح الإجابات'. "
        "وتحقّق من صحة كل إجابة قبل كتابتها، ولا تختلق شواهد ولا مراجع. "
        "نص عادي فقط، بلا markdown وبلا LaTeX وبلا علامات دولار. "
        + AR_ONLY_RULE
    )


SYS_BUILDERS = {
    "rep": sys_report,
    "exp": sys_explain,
    "math": sys_math,
    "code": sys_code,
    "quiz": sys_quiz,
}


# ══════════════════════════════ الواجهة ══════════════════════════════

def main_menu_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🎓 شرح موضوع أو صفحة ملزمة", callback_data="svc:exp")],
        [InlineKeyboardButton("📝 إنشاء تقرير (Word / PDF)", callback_data="svc:rep")],
        [InlineKeyboardButton("🧮 حل مسألة (نص أو صورة)", callback_data="svc:math")],
        [InlineKeyboardButton("💻 تدقيق كود (نص أو صورة)", callback_data="svc:code")],
        [InlineKeyboardButton("📚 بنك أسئلة", callback_data="svc:quiz")],
        [InlineKeyboardButton("⚙️ الإعدادات", callback_data="settings")],
    ])


def subjects_kb(service: str) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for term in (1, 2):
        items = [(c, s) for c, s in SUBJECTS.items()
                 if s["term"] == term and (s["is_code"] if service == "code" else True)]
        if not items:
            continue
        rows.append([InlineKeyboardButton(
            f"— الكورس {'الأول' if term == 1 else 'الثاني'} —", callback_data="noop")])
        for code, sub in items:
            rows.append([InlineKeyboardButton(sub["name"], callback_data=f"sub:{code}")])
    rows.append([InlineKeyboardButton("🔙 القائمة الرئيسية", callback_data="menu")])
    return InlineKeyboardMarkup(rows)


def settings_kb(context: ContextTypes.DEFAULT_TYPE) -> InlineKeyboardMarkup:
    cover = context.user_data.get("cover", False)
    return InlineKeyboardMarkup([
        [InlineKeyboardButton(
            f"صفحة الغلاف: {'مفعّلة ✅' if cover else 'معطّلة ❌'}",
            callback_data="cover:toggle")],
        [InlineKeyboardButton("🔙 القائمة الرئيسية", callback_data="menu")],
    ])


def download_kb(context: ContextTypes.DEFAULT_TYPE) -> InlineKeyboardMarkup:
    cover = (context.user_data or {}).get("cover", False)
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📄 تحميل Word", callback_data="dl:word"),
         InlineKeyboardButton("📕 تحميل PDF", callback_data="dl:pdf")],
        [InlineKeyboardButton(
            f"{'إزالة' if cover else 'إضافة'} صفحة الغلاف", callback_data="cover:toggle")],
        [InlineKeyboardButton("🔙 القائمة الرئيسية", callback_data="menu")],
    ])


# القائمة الرئيسية موجزة: شرح كل خدمة انتقل إلى داخلها، وسطر الأوامر حُذف
WELCOME = (
    "🎓 *المنصة الأكاديمية الذكية*\n"
    "المرحلة الثانية — علوم الحاسوب\n\n"
    "اختر الخدمة، ثم المادة، ثم أرسل طلبك نصاً أو صورة."
)

# وصف كل خدمة يُعرض بعد الضغط على زرها فقط
SERVICE_NOTES = {
    "exp": "تقبل هذه الخدمة اسم الموضوع أو صورة صفحة من الملزمة، "
           "وتقترح لك مقاطع يوتيوب للموضوع.",
    "rep": "تقبل هذه الخدمة عنواناً نصياً، وتُخرج التقرير بصيغة Word أو PDF.",
    "math": "تقبل هذه الخدمة نص المسألة أو صورتها، وتعرض الحل خطوة بخطوة.",
    "code": "تقبل هذه الخدمة الكود نصاً أو صورة، وتعرض الأخطاء والكود المصحّح.\n"
            "(في هذه الخدمة تظهر المواد البرمجية فقط)",
    "quiz": "تقبل هذه الخدمة اسم الموضوع أو الفصل، وتُخرج أسئلة مع مفتاح الإجابات.",
}

# تنبيه يُذيَّل به كل ناتج: النموذج قد يخطئ ولا يعرف محتوى المحاضرة الأصلية
DISCLAIMER = (
    "\n\n⚠️ هذا ناتج آلي مساعد للمراجعة ولا يُغني عن المحاضرة الأصلية. "
    "راجع الصفحة، وتحقّق بنفسك من أي أرقام أو حسابات."
)


# ══════════════════════════════ الأوامر ══════════════════════════════

async def start_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    context.user_data.pop("state", None)
    u = update.effective_user
    await record_activity(u.id, None, u.full_name or "", u.username or "")
    await update.message.reply_text(WELCOME, reply_markup=main_menu_kb(),
                                    parse_mode="Markdown")


async def cancel_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    context.user_data.pop("state", None)
    context.user_data.pop("await_forward", None)
    BUSY_USERS.discard(update.effective_user.id)
    await update.message.reply_text("تم إلغاء الطلب الحالي.", reply_markup=main_menu_kb())


async def settings_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    name = context.user_data.get("student", "غير محدد")
    await update.message.reply_text(
        f"الإعدادات\nاسم الطالب على الغلاف: {name}\n"
        f"السنة الدراسية على الغلاف: {academic_year()}\n"
        "لتغيير الاسم: /name الاسم الثلاثي",
        reply_markup=settings_kb(context),
    )


async def name_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    arg = " ".join(context.args).strip()
    if not arg:
        await update.message.reply_text("اكتب الاسم بعد الأمر، مثال:\n/name أحمد علي حسين")
        return
    context.user_data["student"] = arg[:80]
    await update.message.reply_text(f"تم الحفظ: {arg[:80]}")


async def admin_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if update.effective_user.id != ADMIN_ID:
        return
    st = await get_stats()
    ok, why = _pdf_ready()
    bold_ok = os.path.exists(AR_TTF_BOLD)
    mono_ok = os.path.exists(MONO_TTF)
    yt = "مفتاح API مضبوط ✅" if YOUTUBE_API_KEY else "روابط بحث فقط (بلا مفتاح)"
    await update.message.reply_text(
        "🛠 لوحة تحكم المنصة الأكاديمية\n\n"
        f"الطلبة المسجلين: {len(st['users'])}\n"
        f"مجموع العمليات: {st['requests_count']}\n"
        f"طلبات قيد التنفيذ الآن: {len(BUSY_USERS)}\n\n"
        f"مسار البيانات: {STATS_FILE}\n"
        f"قناة النسخ: {BACKUP_CHAT_ID or 'غير مضبوطة'}\n"
        f"النسخ التلقائي: كل {BACKUP_HOURS} ساعة\n"
        f"السنة الدراسية: {academic_year()}"
        f"{' (مثبّتة يدوياً)' if ACADEMIC_YEAR_ENV else ' (محسوبة تلقائياً)'}\n"
        f"يوتيوب: {yt}\n"
        f"ذاكرة بحث يوتيوب: {len(_YT_CACHE)} عبارة\n"
        f"نتائج لكل بحث: {YT_FETCH} تُفلتر وتُرتّب إلى {YT_RESULTS}\n"
        f"عبارات لكل طلب: {YT_MAX_QUERIES} (كل عبارة = 100 وحدة)\n"
        f"لغة الشرح: {'العربية دائماً' if EXPLAIN_IN_ARABIC else 'حسب لغة المادة'}\n"
        f"PDF: {'جاهز ✅' if ok else 'غير جاهز ❌ — ' + why}\n"
        f"الخط العادي: {'موجود' if os.path.exists(AR_TTF) else 'مفقود'}\n"
        f"الخط العريض: {'موجود' if bold_ok else 'بديل (يستخدم العادي)'}\n"
        f"خط الكود: {'موجود' if mono_ok else 'بديل (يستخدم العادي)'}\n"
        f"محرّك التشكيل: {'uharfbuzz' if FPDF_OK else '-'} / bidi={BIDI_OK}\n"
        f"ضغط الصور: {'مفعّل' if PIL_OK else 'معطّل'}\n\n"
        "الأوامر: /stats /users /backup /restore /getid /broadcast"
    )


async def stats_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if update.effective_user.id != ADMIN_ID:
        return
    st = await get_stats()
    for part in split_for_telegram(build_stats_text(st)):
        await update.message.reply_text(part)


async def users_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if update.effective_user.id != ADMIN_ID:
        return
    st = await get_stats()
    rows = ["id\tname\tusername\trequests\tfirst_seen\tlast_seen"]
    for uid, r in sorted(st["users"].items(), key=lambda x: -x[1].get("requests", 0)):
        rows.append("\t".join([uid, r.get("name", ""), r.get("username", ""),
                               str(r.get("requests", 0)),
                               r.get("first_seen", ""), r.get("last_seen", "")]))
    buf = io.BytesIO("\n".join(rows).encode("utf-8"))
    await update.message.reply_document(document=buf, filename="users.tsv",
                                        caption=f"عدد الطلبة: {len(st['users'])}")


async def backup_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    global _stats_dirty
    if update.effective_user.id != ADMIN_ID:
        return
    if await do_backup(context.bot, "يدوي"):
        _stats_dirty = False
        await update.message.reply_text("تم إرسال النسخة إلى القناة وتثبيتها ✅")
        return
    st = await get_stats()
    blob = json.dumps(st, ensure_ascii=False, indent=1).encode("utf-8")
    await update.message.reply_document(
        document=io.BytesIO(blob), filename="bot_stats.json",
        caption="قناة النسخ غير مضبوطة أو فشل الإرسال، احفظ هذا الملف يدوياً.")


async def restore_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if update.effective_user.id != ADMIN_ID:
        return
    data = await restore_stats(context.bot)
    if data:
        await update.message.reply_text(
            f"تمت الاستعادة ✅\nالطلبة: {len(data['users'])}\n"
            f"العمليات: {data['requests_count']}")
    else:
        await update.message.reply_text(
            "لم أجد نسخة مثبّتة في القناة. يمكنك إرسال ملف bot_stats.json هنا مباشرة.")


async def getid_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if update.effective_user.id != ADMIN_ID:
        return
    context.user_data["await_forward"] = True
    await update.message.reply_text(
        "أعد توجيه أي رسالة من القناة الآن، وسأخبرك بمعرّفها.")


async def _broadcast_worker(bot, ids: list[int], text: str, admin_id: int) -> None:
    sent, failed, blocked = 0, 0, []
    for uid in ids:
        try:
            await bot.send_message(uid, text)
            sent += 1
        except Forbidden:
            blocked.append(uid)
            failed += 1
        except RetryAfter as e:
            await asyncio.sleep(e.retry_after + 1)
            try:
                await bot.send_message(uid, text)
                sent += 1
            except Exception:
                failed += 1
        except Exception:
            failed += 1
        await asyncio.sleep(0.03)
    await remove_users(blocked)
    try:
        await bot.send_message(
            admin_id,
            f"انتهى الإرسال الجماعي.\nنجح: {sent}\nفشل: {failed}\n"
            f"حُذف من القائمة (حاظرون): {len(blocked)}",
        )
    except Exception:
        pass


async def broadcast_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if update.effective_user.id != ADMIN_ID:
        return
    text = " ".join(context.args).strip()
    if not text:
        await update.message.reply_text("اكتب النص بعد الأمر:\n/broadcast مرحباً بالطلبة")
        return
    st = await get_stats()
    ids = [int(u) for u in st["users"].keys() if str(u) != str(ADMIN_ID)]
    await update.message.reply_text(f"بدأ الإرسال إلى {len(ids)} مستخدم…")
    asyncio.create_task(_broadcast_worker(context.bot, ids, text, ADMIN_ID))


# ══════════════════════════════ الأزرار ══════════════════════════════

async def handle_callback(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    q = update.callback_query
    data = q.data or ""
    await q.answer()

    if data == "noop":
        return

    if data == "menu":
        context.user_data.pop("state", None)
        await q.edit_message_text(WELCOME, reply_markup=main_menu_kb(),
                                  parse_mode="Markdown")
        return

    if data == "settings":
        await q.edit_message_text(
            "الإعدادات\n\nصفحة الغلاف تُضاف للتقرير أو الشرح فقط عند تفعيلها، "
            "والوضع الافتراضي هو النص وحده.\n"
            f"السنة الدراسية الحالية على الغلاف: {academic_year()}\n"
            "لتعيين اسمك على الغلاف: /name الاسم الثلاثي",
            reply_markup=settings_kb(context),
        )
        return

    if data == "cover:toggle":
        context.user_data["cover"] = not context.user_data.get("cover", False)
        on = context.user_data["cover"]
        if context.user_data.get("last_report"):
            await q.edit_message_text(
                f"صفحة الغلاف أصبحت {'مفعّلة ✅' if on else 'معطّلة ❌'}.\n"
                "اختر صيغة التحميل:",
                reply_markup=download_kb(context),
            )
        else:
            await q.edit_message_text(
                f"صفحة الغلاف أصبحت {'مفعّلة ✅' if on else 'معطّلة ❌'}.",
                reply_markup=settings_kb(context),
            )
        return

    if data.startswith("svc:"):
        svc = data.split(":", 1)[1]
        if svc not in SYS_BUILDERS:
            return
        context.user_data["state"] = {"service": svc}
        note = SERVICE_NOTES.get(svc, "")
        await q.edit_message_text(
            f"الخدمة: {SERVICE_LABELS.get(svc, svc)}\n"
            f"{note}\n\nاختر المادة:\n\nللإلغاء في أي وقت: /cancel",
            reply_markup=subjects_kb(svc),
        )
        return

    if data.startswith("sub:"):
        code = data.split(":", 1)[1]
        sub = SUBJECTS.get(code)
        state = context.user_data.get("state") or {}
        if not sub or "service" not in state:
            await q.edit_message_text("انتهت صلاحية الاختيار. ابدأ من جديد.",
                                      reply_markup=main_menu_kb())
            return
        state.update({"code": code, "subject": sub["name"], "lang": sub["lang"]})
        context.user_data["state"] = state
        svc = state["service"]
        prompts = {
            "rep": "أرسل عنوان التقرير المطلوب.",
            "exp": ("أرسل اسم الموضوع الذي تريد شرحه،\n"
                    "أو صوّر صفحة من الملزمة وأرسل الصورة (حتى 5 صور كألبوم).\n"
                    "يمكنك إضافة تعليق مع الصورة مثل: «اشرح لي الجزء الثاني فقط»."),
            "math": "أرسل نص المسألة، أو صوّرها وأرسل الصورة (حتى 5 صور كألبوم).",
            "code": ("أرسل الكود نصاً (الأفضل والأدق)، أو صورة له.\n"
                     "ملاحظة: تصوير الكود من محرر بواجهة عربية قد يقلب مواضع الأقواس."),
            "quiz": "أرسل الموضوع أو الفصل المطلوب إنشاء أسئلة عنه.",
        }
        if svc == "exp":
            lang_note = ("سيكون الشرح بالعربية مع الحفاظ على المصطلحات الإنكليزية"
                         if explain_lang(sub) == "ar" and sub["lang"] == "en"
                         else ("سيكون الشرح بالعربية" if explain_lang(sub) == "ar"
                               else "سيكون الشرح باللغة الإنكليزية"))
        else:
            lang_note = ("سيكون الناتج باللغة الإنكليزية" if sub["lang"] == "en"
                         else "سيكون الناتج باللغة العربية")
        await q.edit_message_text(
            f"المادة: {sub['name']}\n{lang_note}\n\n"
            f"{prompts[svc]}\n\nللإلغاء: /cancel"
        )
        return

    if data.startswith("dl:"):
        fmt = data.split(":", 1)[1]
        rep = context.user_data.get("last_report")
        if not rep:
            await q.edit_message_text("لا يوجد ملف محفوظ. أنشئ تقريراً أو شرحاً جديداً.",
                                      reply_markup=main_menu_kb())
            return
        cover = context.user_data.get("cover", False)
        student = context.user_data.get("student", "")
        base = safe_filename(rep["title"])
        chat_id = q.message.chat_id
        try:
            if fmt == "word":
                await context.bot.send_chat_action(chat_id, ChatAction.UPLOAD_DOCUMENT)
                blob = await asyncio.to_thread(
                    create_academic_docx, rep["title"], rep["body"], rep["subject"],
                    rep["lang"], cover, student)
                await context.bot.send_document(
                    chat_id, document=io.BytesIO(blob), filename=f"{base}.docx",
                    caption=f"📄 {rep['title']}")
                await record_file("word")
            else:
                ok, why = _pdf_ready(rep["lang"])
                if not ok:
                    await context.bot.send_message(chat_id, f"تعذّر إنشاء PDF: {why}")
                    return
                await context.bot.send_chat_action(chat_id, ChatAction.UPLOAD_DOCUMENT)
                blob = await asyncio.to_thread(
                    create_academic_pdf, rep["title"], rep["body"], rep["subject"],
                    rep["lang"], cover, student)
                await context.bot.send_document(
                    chat_id, document=io.BytesIO(blob), filename=f"{base}.pdf",
                    caption=f"📕 {rep['title']}")
                await record_file("pdf")
        except Exception as e:
            logger.exception("download failed")
            if ADMIN_ID and q.from_user.id == ADMIN_ID:
                await context.bot.send_message(
                    chat_id, f"تفاصيل الخطأ:\n{type(e).__name__}: {e}"[:3500])
            else:
                await context.bot.send_message(
                    chat_id, "حدث خطأ أثناء توليد الملف. جرّب الصيغة الأخرى.")
        return


# ══════════════════════════════ معالجة النص ══════════════════════════════

async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user = update.effective_user
    msg = update.message
    text = (msg.text or "").strip()
    if not text:
        return

    state = context.user_data.get("state")
    if not state or "code" not in state:
        await msg.reply_text("اختر الخدمة والمادة أولاً:", reply_markup=main_menu_kb())
        return

    if user.id in BUSY_USERS:
        await msg.reply_text("طلبك السابق قيد المعالجة، انتظر انتهاءه أو أرسل /cancel")
        return

    if len(text) > MAX_INPUT_CHARS:
        await msg.reply_text(f"النص طويل جداً ({len(text)} حرف). الحد {MAX_INPUT_CHARS} حرف.")
        return

    svc = state["service"]
    sub = SUBJECTS[state["code"]]
    lang = explain_lang(sub) if svc == "exp" else sub["lang"]

    BUSY_USERS.add(user.id)
    notice = await msg.reply_text("⏳ جارٍ المعالجة…")
    try:
        await context.bot.send_chat_action(msg.chat_id, ChatAction.TYPING)
        system = SYS_BUILDERS[svc](sub)
        merge = svc in ("math", "quiz")

        if svc == "rep":
            user_msg = (f"Report title: {text}" if sub["lang"] == "en"
                        else f"عنوان التقرير: {text}")
        elif svc == "exp":
            user_msg = (f"الموضوع المطلوب شرحه: {text}\n"
                        f"المادة: {sub['name']} — {sub.get('en', '')}")
        else:
            user_msg = text

        raw = await query_ai_text(system, user_msg, merge_system=merge,
                                  temperature=0.35 if svc == "exp" else 0.3)

        if svc == "code":
            body = clean_code_output(raw)
        elif svc in ("rep", "exp"):
            body = clean_report_text(raw, lang)
            if lang == "ar" and (svc == "rep" or sub["lang"] == "ar"):
                body = strip_latin_parentheticals(body)
        else:
            body = clean_text_strictly(raw, lang)
            if lang == "ar":
                body = strip_latin_parentheticals(body)

        if not body:
            raise RuntimeError("ناتج فارغ بعد التنظيف")

        await notice.delete()

        if svc == "rep":
            context.user_data["last_report"] = {
                "title": text[:120], "body": body,
                "subject": sub["name"], "lang": lang,
            }
            cover = context.user_data.get("cover", False)
            await msg.reply_text(
                f"✅ تم إنشاء التقرير: {text[:120]}\n"
                f"المادة: {sub['name']}\n"
                f"عدد الكلمات ≈ {len(body.split())}\n"
                f"صفحة الغلاف: {'مفعّلة' if cover else 'معطّلة'}\n\n"
                "اختر صيغة التحميل:" + DISCLAIMER,
                reply_markup=download_kb(context),
            )

        elif svc == "exp":
            # الطالب كتب الموضوع بنفسه، فلا يُستخرج من المتن
            # (كان يُعاد «نظرة عامة» فتفسد عبارة بحث يوتيوب واسم الملف)
            topic = text[:120]
            context.user_data["last_report"] = {
                "title": f"شرح: {topic}"[:120], "body": body,
                "subject": sub["name"], "lang": lang,
            }
            for part in split_for_telegram(headings_for_telegram(body)):
                await msg.reply_text(part)
            try:
                yt = await build_youtube_block(topic, sub)
                if yt:
                    await send_plain(context.bot, msg.chat_id, yt, parse_mode="HTML")
            except Exception as e:
                logger.warning("youtube block failed: %s", e)
            await msg.reply_text(
                "يمكنك تحميل الشرح كملف للمراجعة، أو إرسال موضوع آخر لشرحه:"
                + DISCLAIMER,
                reply_markup=download_kb(context),
            )

        else:
            for part in split_for_telegram(body):
                await msg.reply_text(part)
            await msg.reply_text("للطلب التالي اختر من القائمة:" + DISCLAIMER,
                                 reply_markup=main_menu_kb())

        await record_activity(user.id, svc, user.full_name or "",
                              user.username or "", state["code"])

    except Exception:
        logger.exception("handle_text failed")
        try:
            await notice.edit_text("تعذّر إكمال الطلب حالياً. أعد المحاولة بعد قليل.")
        except Exception:
            await msg.reply_text("تعذّر إكمال الطلب حالياً. أعد المحاولة بعد قليل.")
    finally:
        BUSY_USERS.discard(user.id)
        if svc not in ("rep", "exp"):
            context.user_data.pop("state", None)


# ══════════════════════════════ معالجة الصور ══════════════════════════════

async def run_image_request(context: ContextTypes.DEFAULT_TYPE, chat_id: int, user_id: int,
                            state: dict, file_ids: list[str], caption: str,
                            name: str = "", username: str = "") -> None:
    svc = state["service"]
    sub = SUBJECTS[state["code"]]
    lang = explain_lang(sub) if svc == "exp" else sub["lang"]

    if user_id in BUSY_USERS:
        await context.bot.send_message(chat_id, "طلبك السابق قيد المعالجة، انتظر انتهاءه.")
        return

    BUSY_USERS.add(user_id)
    wait_note = ("🔍 جارٍ قراءة الصفحة وتحضير الشرح…" if svc == "exp"
                 else "🔍 جارٍ قراءة الصورة ومعالجتها…")
    notice = await context.bot.send_message(chat_id, wait_note)
    try:
        await context.bot.send_chat_action(chat_id, ChatAction.TYPING)
        images = await fetch_images(context.bot, file_ids)
        if not images:
            raise RuntimeError("لا توجد صور صالحة")

        parts = [SYS_BUILDERS[svc](sub), ""]
        parts.append(OCR_EXPLAIN if svc == "exp" else OCR_PREFIX)
        if svc == "code":
            parts.append(OCR_CODE_EXTRA)
        if caption:
            parts.append(("تعليق الطالب: " if lang == "ar" else "Student note: ") + caption)
        prompt = "\n\n".join(parts)

        raw = await query_ai_vision(prompt, images)

        if svc == "code":
            body = clean_code_output(raw)
        elif svc == "exp":
            body = clean_report_text(raw, lang)
            if lang == "ar" and sub["lang"] == "ar":
                body = strip_latin_parentheticals(body)
        else:
            body = clean_text_strictly(raw, lang)
            if lang == "ar":
                body = strip_latin_parentheticals(body)

        if not body:
            raise RuntimeError("ناتج فارغ")

        await notice.delete()

        if svc == "exp":
            topic = extract_topic(body, caption or sub["name"])
            kb = main_menu_kb()
            try:
                if context.user_data is not None:
                    context.user_data["last_report"] = {
                        "title": f"شرح: {topic}"[:120], "body": body,
                        "subject": sub["name"], "lang": lang,
                    }
                    kb = download_kb(context)
            except Exception as e:
                logger.warning("cannot store explanation: %s", e)

            for part in split_for_telegram(headings_for_telegram(body)):
                await context.bot.send_message(chat_id, part)
            try:
                yt = await build_youtube_block(topic, sub)
                if yt:
                    await send_plain(context.bot, chat_id, yt, parse_mode="HTML")
            except Exception as e:
                logger.warning("youtube block failed: %s", e)
            await context.bot.send_message(
                chat_id,
                "يمكنك تحميل الشرح كملف، أو إرسال صفحة أخرى لشرحها:" + DISCLAIMER,
                reply_markup=kb,
            )
        else:
            for part in split_for_telegram(body):
                await context.bot.send_message(chat_id, part)
            await context.bot.send_message(chat_id,
                                           "للطلب التالي اختر من القائمة:" + DISCLAIMER,
                                           reply_markup=main_menu_kb())

        await record_activity(user_id, f"{svc}_img", name, username, state["code"])

    except VisionUnavailable:
        # نماذج الرؤية على Groq من فئة المعاينة وقد تُسحب بإشعار قصير
        logger.exception("vision chain unavailable")
        txt = ("خدمة قراءة الصور غير متاحة الآن. "
               "أرسل السؤال أو الكود مكتوباً نصاً وسيعمل الطلب عادياً.")
        try:
            await notice.edit_text(txt)
        except Exception:
            await context.bot.send_message(chat_id, txt)
    except Exception:
        logger.exception("image request failed")
        try:
            await notice.edit_text("تعذّرت قراءة الصورة. جرّب صورة أوضح أو أرسل النص مكتوباً.")
        except Exception:
            await context.bot.send_message(chat_id, "تعذّرت قراءة الصورة. جرّب صورة أوضح.")
    finally:
        BUSY_USERS.discard(user_id)


async def _album_job(context: ContextTypes.DEFAULT_TYPE) -> None:
    job = context.job
    key = job.data["key"]
    state = job.data["state"]
    buf = (context.chat_data or {}).pop(key, None)
    if not buf:
        return
    await run_image_request(context, job.chat_id, job.user_id, state,
                            buf["file_ids"], buf.get("caption", ""),
                            job.data.get("name", ""), job.data.get("username", ""))


async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg = update.message
    user = update.effective_user
    state = context.user_data.get("state")

    if not state or "code" not in state:
        await msg.reply_text("اختر الخدمة والمادة أولاً ثم أرسل الصورة:",
                             reply_markup=main_menu_kb())
        return

    svc = state["service"]
    if svc == "rep":
        await msg.reply_text("خدمة التقرير تحتاج عنواناً نصياً. أرسل عنوان التقرير.")
        return
    if svc == "quiz":
        await msg.reply_text("خدمة بنك الأسئلة تحتاج اسم الموضوع نصاً.")
        return

    fid = msg.photo[-1].file_id
    caption = (msg.caption or "").strip()

    if msg.media_group_id:
        key = f"album:{msg.media_group_id}"
        buf = context.chat_data.setdefault(key, {"file_ids": [], "caption": ""})
        if len(buf["file_ids"]) < MAX_IMAGES:
            buf["file_ids"].append(fid)
        if caption and not buf["caption"]:
            buf["caption"] = caption
        if context.job_queue:
            for j in context.job_queue.get_jobs_by_name(key):
                j.schedule_removal()
            context.job_queue.run_once(
                _album_job, ALBUM_WAIT, name=key,
                chat_id=msg.chat_id, user_id=user.id,
                data={"key": key, "state": state,
                      "name": user.full_name or "", "username": user.username or ""},
            )
        return

    await run_image_request(context, msg.chat_id, user.id, state, [fid], caption,
                            user.full_name or "", user.username or "")


async def handle_forwarded(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """يعمل بعد /getid فقط؛ وإلا يمرّر الرسالة إلى معالجها الطبيعي."""
    msg = update.message
    if not context.user_data.pop("await_forward", False):
        if msg.photo:
            await handle_photo(update, context)
        elif msg.text:
            await handle_text(update, context)
        return
    origin = getattr(msg, "forward_origin", None)
    src = getattr(origin, "chat", None) or getattr(origin, "sender_chat", None)
    lines = [f"معرّف هذه الدردشة: {msg.chat_id}"]
    if src is not None:
        lines += [f"نوع المصدر: {src.type}",
                  f"العنوان: {src.title or '-'}",
                  f"معرّف المصدر: {src.id}", "",
                  "ضع معرّف المصدر في المتغيّر BACKUP_CHAT_ID"]
    else:
        lines.append("تعذّر قراءة معرّف المصدر. أعد التوجيه من القناة نفسها.")
    await msg.reply_text("\n".join(lines))


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg = update.message
    doc = msg.document
    if (update.effective_user.id == ADMIN_ID and doc
            and (doc.file_name or "").lower().endswith(".json")):
        try:
            f = await doc.get_file()
            raw = bytes(await f.download_as_bytearray())
            incoming = json.loads(raw.decode("utf-8"))
            async with _stats_lock:
                data = await asyncio.to_thread(_read_stats_file)
                data = _merge_stats(data, incoming)
                await asyncio.to_thread(_write_stats_file, data)
            await msg.reply_text(
                f"تم دمج النسخة ✅\nالطلبة: {len(data['users'])}\n"
                f"العمليات: {data['requests_count']}")
        except Exception as e:
            await msg.reply_text(f"فشل قراءة الملف: {e}")
        return
    await msg.reply_text("الملفات غير مدعومة. أرسل النص مكتوباً أو صورة للصفحة أو المسألة.")


async def error_handler(update: object, context: ContextTypes.DEFAULT_TYPE) -> None:
    logger.error("unhandled error", exc_info=context.error)


# ══════════════════════════════ خادم الفحص ══════════════════════════════

class SimpleHealthHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        body = b'{"status":"ok","service":"academic_bot"}'
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_HEAD(self):
        self.send_response(200)
        self.end_headers()

    def log_message(self, *args):
        return


def run_http_server() -> None:
    try:
        ThreadingHTTPServer(("0.0.0.0", PORT), SimpleHealthHandler).serve_forever()
    except Exception as e:
        logger.error("health server stopped: %s", e)


# ══════════════════════════════ الإقلاع والإيقاف ══════════════════════════════

async def post_init(app) -> None:
    data = await get_stats()
    if not data["users"]:
        restored = await restore_stats(app.bot)
        if restored:
            data = restored
    if app.job_queue and BACKUP_CHAT_ID:
        app.job_queue.run_repeating(backup_job, interval=int(BACKUP_HOURS * 3600),
                                    first=180, name="stats_backup")
        logger.info("auto backup every %s hours", BACKUP_HOURS)
    if ADMIN_ID:
        try:
            await app.bot.send_message(
                ADMIN_ID,
                f"البوت يعمل الآن ✅\nالطلبة: {len(data['users'])}\n"
                f"العمليات: {data['requests_count']}\n"
                f"السنة الدراسية: {academic_year()}\n"
                f"يوتيوب: {'مفتاح مضبوط' if YOUTUBE_API_KEY else 'روابط بحث فقط'}")
        except Exception:
            pass


async def post_stop(app) -> None:
    await do_backup(app.bot, "إيقاف")


# ══════════════════════════════ التشغيل ══════════════════════════════

def main() -> None:
    if not TELEGRAM_BOT_TOKEN:
        raise SystemExit("TELEGRAM_BOT_TOKEN غير موجود في متغيرات البيئة.")
    if not GROQ_API_KEY:
        raise SystemExit("GROQ_API_KEY غير موجود في متغيرات البيئة.")

    os.makedirs(DATA_DIR, exist_ok=True)
    threading.Thread(target=run_http_server, daemon=True).start()
    logger.info("health server on port %s", PORT)
    logger.info("PDF: fpdf2=%s font=%s bold=%s mono=%s bidi_fallback=%s",
                FPDF_OK, os.path.exists(AR_TTF), os.path.exists(AR_TTF_BOLD),
                os.path.exists(MONO_TTF), BIDI_OK)
    logger.info("stats file: %s | backup chat: %s", STATS_FILE, BACKUP_CHAT_ID or "none")
    logger.info("academic year: %s | youtube key: %s | queries per request: %s",
                academic_year(), bool(YOUTUBE_API_KEY), YT_MAX_QUERIES)

    persistence = PicklePersistence(filepath=PERSIST_FILE)
    app = (
        Application.builder()
        .token(TELEGRAM_BOT_TOKEN)
        .persistence(persistence)
        .post_init(post_init)
        .post_stop(post_stop)
        .concurrent_updates(True)
        .build()
    )

    app.add_handler(CommandHandler("start", start_command))
    app.add_handler(CommandHandler("cancel", cancel_command))
    app.add_handler(CommandHandler("settings", settings_command))
    app.add_handler(CommandHandler("name", name_command))
    app.add_handler(CommandHandler("admin", admin_command))
    app.add_handler(CommandHandler("stats", stats_command))
    app.add_handler(CommandHandler("users", users_command))
    app.add_handler(CommandHandler("backup", backup_command))
    app.add_handler(CommandHandler("restore", restore_command))
    app.add_handler(CommandHandler("getid", getid_command))
    app.add_handler(CommandHandler("broadcast", broadcast_command))
    app.add_handler(CallbackQueryHandler(handle_callback))
    if ADMIN_ID:
        app.add_handler(MessageHandler(
            filters.FORWARDED & filters.User(ADMIN_ID), handle_forwarded))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_error_handler(error_handler)

    logger.info("bot started")
    app.run_polling(drop_pending_updates=True, allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
